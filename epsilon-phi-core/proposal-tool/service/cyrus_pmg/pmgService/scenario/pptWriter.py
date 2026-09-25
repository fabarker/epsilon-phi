"""The deck renderer - the SheetDocs as PowerPoint, natively (D121).

Phase 3 of the PowerPoint export plan (proposals/powerpoint-export-plan.html):
``writeDeck`` renders the same Docs the Excel renderer transcribes - built by
``sheetDoc``'s builders, paginated by ``sheetDoc.paginate`` - as one slide per
page: a heading, a subheading, the table in the report's own faces, fills and
rules, footnotes, and the Proposal UID in the footer and the file's core
properties (``dc:identifier``, so ``stampedProposalId`` reads a deck exactly
as it reads a workbook).

Like ``workbook.py``, this module decides nothing about the report: every
colour, face, weight, border and number format comes from the Doc, and the
number formats are rendered to text by ``sheetDoc.renderNumber`` because a
PowerPoint table cell holds text. The two helpers that write raw XML exist
because python-pptx has no first-class API for table cell borders or for a
doughnut's hole size; each writes only what the Doc or the chart spec says.

The stress block's red/green is the Doc's sign rule resolved per cell from
the RAW value with strict inequalities - a -0.04% that prints as -0.0% is
red, an exact zero keeps the base ink - which is precisely how the Excel
renderer's CellIsRule behaves.

Scale is the workbook's own print convention transplanted: a page's table
scales to the slide's content box by ONE factor - fonts, row heights and
column widths together, ``fitToWidth=1`` for a fixed canvas - so proportions
survive even where absolute sizes cannot.
"""

from __future__ import annotations

import datetime
import io

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from . import sheetDoc
from .sheetDoc import SheetDoc, paginate, renderNumber
from .workbook import (DONUT_DIMENSIONS, DONUT_PALETTE,
                       buildImplementationRows, donutBreakdown, implColumns,
                       _TEXT_COLUMNS, _WIDTHS)

# ---- geometry: one 16:9 canvas, one content box --------------------------
SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5
MARGIN_IN = 0.45
CONTENT_TOP_IN = 1.16                 # below heading, subheading and rule
FOOTER_H_IN = 0.62
CONTENT_W_PT = (SLIDE_W_IN - 2 * MARGIN_IN) * 72.0
CONTENT_H_PT = (SLIDE_H_IN - CONTENT_TOP_IN - FOOTER_H_IN) * 72.0
#: the pagination budget: enough rows that the scale floor stays legible
#: (CONTENT_H / 730 ~ 0.56, so 12pt table text never drops below ~6.8pt)
PAGE_BUDGET_PT = 730.0
#: one Excel column-width character, in points (the width of '0' at 96dpi)
CHAR_PT = 5.25
DEFAULT_COL_CHARS = 8.43              # Excel's default column width
#: the deck's own chrome inks - slide furniture, not report styling
_HEAD_INK = '0F243E'
_SUB_INK = '4B5A6A'
_NOTE_INK = '7C8B9C'
#: what Excel shows for a cell no style touched
_DEFAULT_FONT = {'name': 'Calibri', 'size': 11}
#: border kinds to (width in points, dash) - the Excel edges, drawn
_EDGE_PT = {'hair': (0.25, None), 'dotted': (0.75, 'sysDot'),
            'thin': (1.0, None), 'thick': (2.25, None)}

_A_NS = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'


def _in(inches: float) -> Emu:
    return Emu(int(round(inches * 914400)))


def _pt(points: float) -> Emu:
    return Emu(int(round(points * 12700)))


# ---- the two XML helpers -------------------------------------------------

def _edgeXml(tag: str, edge, scale: float) -> str:
    """One a:lnX element: the Doc's edge, or an explicit no-line - because a
    PowerPoint table style draws borders of its own, and the only way to get
    exactly the Doc's rules is to state every edge."""
    if edge is None:
        return '<a:{tag} {ns}><a:noFill/></a:{tag}>'.format(tag=tag, ns=_A_NS)
    style, colour = edge
    widthPt, dash = _EDGE_PT[style]
    return ('<a:{tag} {ns} w="{w}" cap="flat"><a:solidFill>'
            '<a:srgbClr val="{c}"/></a:solidFill>{dash}</a:{tag}>').format(
                tag=tag, ns=_A_NS, w=int(round(widthPt * scale * 12700)),
                c=colour or '000000',
                dash='<a:prstDash val="{}"/>'.format(dash) if dash else '')


def _setCellBorders(cell, border, scale: float) -> None:
    """Write all four edges of a table cell. The Doc only ever draws top and
    bottom rules; left and right are stated as no-line so nothing inherited
    from the table style survives."""
    border = border or {}
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
        for element in tcPr.findall(qn(tag)):
            tcPr.remove(element)
    ordered = [('lnL', border.get('left')), ('lnR', border.get('right')),
               ('lnT', border.get('top')), ('lnB', border.get('bottom'))]
    for position, (tag, edge) in enumerate(ordered):
        tcPr.insert(position, parse_xml(_edgeXml(tag, edge, scale)))


def _setHoleSize(chart, value: int) -> None:
    """The doughnut's hole, matched to the workbook's 55 - python-pptx has
    no property for it, so the one element is written by hand."""
    space = chart._chartSpace
    holes = space.findall('.//' + qn('c:holeSize'))
    if holes:
        for hole in holes:
            hole.set('val', str(value))
        return
    for plot in space.findall('.//' + qn('c:doughnutChart')):
        element = plot.makeelement(qn('c:holeSize'), {'val': str(value)})
        plot.append(element)


# ---- slide chrome --------------------------------------------------------

def _textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(_in(left), _in(top), _in(width), _in(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    return frame


def _para(frame, text, sizePt, colour, bold=False, first=False,
          align=PP_ALIGN.LEFT, name='Calibri Light'):
    paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = name
    run.font.size = Pt(sizePt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(colour)
    return paragraph


def _chrome(slide, heading, subheading, footnotes, footer, pageText):
    """Heading, subheading, the navy rule, footnotes bottom-left and the
    identity bottom-right - the frame every content slide shares."""
    head = _textbox(slide, MARGIN_IN, 0.26, SLIDE_W_IN - 2 * MARGIN_IN, 0.44)
    _para(head, heading, 22, _HEAD_INK, bold=True, first=True)
    if subheading:
        sub = _textbox(slide, MARGIN_IN, 0.70, SLIDE_W_IN - 2 * MARGIN_IN, 0.34)
        _para(sub, subheading, 11, _SUB_INK, first=True)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, _in(MARGIN_IN), _in(1.05),
                                  _in(SLIDE_W_IN - 2 * MARGIN_IN), _pt(2.5))
    rule.fill.solid()
    rule.fill.fore_color.rgb = RGBColor.from_string(_HEAD_INK)
    rule.line.fill.background()
    rule.shadow.inherit = False
    if footnotes:
        notes = _textbox(slide, MARGIN_IN, SLIDE_H_IN - 0.58,
                         SLIDE_W_IN - 2 * MARGIN_IN - 2.6, 0.5)
        for index, note in enumerate(footnotes):
            _para(notes, note, 8, _NOTE_INK, first=index == 0)
    page = _textbox(slide, SLIDE_W_IN - MARGIN_IN - 2.5, SLIDE_H_IN - 0.58,
                    2.5, 0.5)
    _para(page, footer, 8, _NOTE_INK, first=True, align=PP_ALIGN.RIGHT)
    _para(page, pageText, 8, _NOTE_INK, align=PP_ALIGN.RIGHT)


# ---- the table -----------------------------------------------------------

def _docColumns(doc: SheetDoc) -> int:
    return max((max(row.cells) for row in doc.rows.values() if row.cells),
               default=1)


def _columnChars(doc: SheetDoc, columns: int):
    return [doc.widths.get(sheetDoc.columnLetter(index), DEFAULT_COL_CHARS)
            for index in range(1, columns + 1)]


_ALIGN = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER,
          'right': PP_ALIGN.RIGHT}


def _signInk(doc: SheetDoc, rowIndex: int, column: int, value):
    """The sign rule, resolved per cell from the RAW value - strict
    inequalities, so an exact zero keeps the base ink (D121)."""
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        return None
    for col1, row1, col2, row2, negInk, posInk in doc.signRules:
        if row1 <= rowIndex <= row2 and col1 <= column <= col2:
            if value < 0:
                return negInk
            if value > 0:
                return posInk
    return None


def _fillCell(cell, spec, doc, rowIndex, column, scale):
    font = spec.font or _DEFAULT_FONT
    ink = _signInk(doc, rowIndex, column, spec.value) or font.get('color')
    align = spec.align or {}
    frame = cell.text_frame
    frame.word_wrap = bool(align.get('wrap'))
    cell.vertical_anchor = (MSO_ANCHOR.MIDDLE if align.get('vertical') == 'center'
                            else MSO_ANCHOR.BOTTOM)
    pad = _pt(2.0 * scale)
    cell.margin_left = cell.margin_right = pad
    cell.margin_top = cell.margin_bottom = _pt(0.6 * scale)
    # Excel's indent leans on the side the alignment leans on
    indent = align.get('indent') or 0
    if indent:
        extra = _pt(indent * 6.0 * scale)
        if align.get('horizontal') == 'right':
            cell.margin_right = Emu(int(pad) + int(extra))
        else:
            cell.margin_left = Emu(int(pad) + int(extra))
    text = renderNumber(spec.value, spec.fmt)
    alignment = _ALIGN.get(align.get('horizontal'), PP_ALIGN.LEFT)
    for index, line in enumerate(text.split('\n') if text else ['']):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = alignment
        run = paragraph.add_run()
        run.text = line.strip() if index else line
        run.font.name = font.get('name', 'Calibri')
        run.font.size = Pt(max(font.get('size', 11) * scale, 1))
        run.font.bold = bool(font.get('bold'))
        run.font.color.rgb = RGBColor.from_string(ink or '000000')
    if spec.fill:
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string(spec.fill)
    else:
        cell.fill.background()
    _setCellBorders(cell, spec.border, scale)


def _renderTable(slide, doc: SheetDoc, page) -> None:
    """One page of a Doc as a native table, scaled to the content box by one
    factor - fonts, heights and widths together (D121)."""
    rows = page.header + page.body
    columns = _docColumns(doc)
    chars = _columnChars(doc, columns)
    naturalW = sum(chars) * CHAR_PT
    naturalH = sum((doc.rows[r].height if r in doc.rows and
                    doc.rows[r].height is not None else sheetDoc.DEFAULT_ROW_PT)
                   for r in rows)
    scale = min(1.0, CONTENT_W_PT / naturalW, CONTENT_H_PT / naturalH)
    widthIn = naturalW * scale / 72.0
    left = (SLIDE_W_IN - widthIn) / 2.0

    frame = slide.shapes.add_table(len(rows), columns, _in(left),
                                   _in(CONTENT_TOP_IN), _pt(naturalW * scale),
                                   _pt(naturalH * scale))
    table = frame.table
    table.first_row = False
    table.last_row = False
    table.first_col = False
    table.last_col = False
    table.horz_banding = False
    table.vert_banding = False
    for index, width in enumerate(chars):
        table.columns[index].width = _pt(width * CHAR_PT * scale)
    position = {}
    for tableRow, rowIndex in enumerate(rows):
        position[rowIndex] = tableRow
        spec = doc.rows.get(rowIndex)
        height = (spec.height if spec is not None and spec.height is not None
                  else sheetDoc.DEFAULT_ROW_PT)
        table.rows[tableRow].height = _pt(height * scale)
        cells = spec.cells if spec is not None else {}
        for column in range(1, columns + 1):
            _fillCell(table.cell(tableRow, column - 1),
                      cells.get(column, sheetDoc.Cell()),
                      doc, rowIndex, column, scale)
    for row1, col1, row2, col2 in doc.merges:
        if row1 in position and row2 in position:
            table.cell(position[row1], col1 - 1).merge(
                table.cell(position[row2], col2 - 1))


# ---- the charts ----------------------------------------------------------

def _renderDonuts(slide, items) -> None:
    """The five composition doughnuts, native charts fed straight from
    ``donutBreakdown`` - no hidden sheet, because a PowerPoint chart carries
    its data - each slice coloured by its palette slot, hole at 55."""
    across = len(DONUT_DIMENSIONS)
    gap = 0.18
    width = (SLIDE_W_IN - 2 * MARGIN_IN - gap * (across - 1)) / across
    top = CONTENT_TOP_IN + 0.32
    height = SLIDE_H_IN - top - FOOTER_H_IN - 0.25
    for index, (key, label) in enumerate(DONUT_DIMENSIONS):
        slices = donutBreakdown(items, key)
        if not slices:
            continue
        data = CategoryChartData()
        data.categories = [entry['name'] for entry in slices]
        data.add_series('Share', [entry['pct'] for entry in slices])
        left = MARGIN_IN + index * (width + gap)
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.DOUGHNUT, _in(left), _in(top), _in(width),
            _in(height), data).chart
        chart.has_title = True
        chart.chart_title.text_frame.text = label
        title = chart.chart_title.text_frame.paragraphs[0].runs[0].font
        title.size = Pt(11)
        title.bold = True
        title.name = 'Calibri'
        title.color.rgb = RGBColor.from_string(_HEAD_INK)
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(7.5)
        chart.plots[0].has_data_labels = False
        series = chart.series[0]
        for offset, entry in enumerate(slices):
            point = series.points[offset]
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = RGBColor.from_string(
                DONUT_PALETTE[entry['slot']])
            point.format.line.color.rgb = RGBColor.from_string('FFFFFF')
            point.format.line.width = Pt(0.75)
        _setHoleSize(chart, 55)


# ---- the deck ------------------------------------------------------------

def _tableSlice(doc: SheetDoc, fromRow: int) -> SheetDoc:
    """The Doc from *fromRow* on, sharing its cells - the implementation
    sheet's preamble becomes the slide's subheading, not table rows."""
    view = SheetDoc(doc.name)
    view.rows = {index: row for index, row in doc.rows.items()
                 if index >= fromRow}
    view.merges = [m for m in doc.merges if m[0] >= fromRow]
    view.widths = doc.widths
    view.signRules = list(doc.signRules)
    view.sections = {index for index in doc.sections if index >= fromRow}
    return view


def _preamblePairs(doc: SheetDoc, headerRow: int):
    pairs = []
    for index in range(1, headerRow):
        row = doc.rows.get(index)
        if not row or 1 not in row.cells or row.cells[1].value is None:
            continue
        label = str(row.cells[1].value)
        value = row.cells.get(2).value if 2 in row.cells else None
        pairs.append('{} {}'.format(label, value) if value is not None else label)
    return pairs


def _continued(subheading, part, parts):
    if parts == 1:
        return subheading
    return '{} — continued, page {} of {}'.format(subheading, part, parts) \
        if part > 1 else '{} — page {} of {}'.format(subheading, part, parts)


def writeDeck(basis, mandate, results, sleevesMap, autoCategories,
              variant: str = None, tacticalTilt: bool = False,
              feeSchedule: str = None, feeLevel: str = None,
              includeFees: bool = True, volPremium: bool = False,
              assets=None, model: dict = None, proposalId: str = None,
              customFees: dict = None, customFeesBy: str = None,
              customFeesAt=None, cover: bool = True) -> bytes:
    """The proposal deck: the workbook's sheets as slides, same engine (D121).

    Same inputs as ``writeWorkbook``, and the model built once the same way
    (D69). *proposalId* is the delivered proposal the deck cites - in every
    footer and in ``dc:identifier``; absent, every footer says the deck is a
    draft, because a deck with no delivered workbook behind it is not a
    delivery (plan decision 2).
    """
    if not includeFees:
        feeSchedule = None
    if model is None:
        model = buildImplementationRows(results[0], sleevesMap, autoCategories,
                                        mandate.mandateSize, variant,
                                        tacticalTilt, feeSchedule, feeLevel,
                                        mandate.topAccountSize, volPremium,
                                        basis.currency, customFees)

    portDoc = sheetDoc.buildPortfoliosDoc(results)
    riskDoc = sheetDoc.buildRiskDoc(results)
    assumDoc = sheetDoc.buildAssumptionsDoc(assets, results)
    implDoc, headerRow, _ = sheetDoc.buildImplementationDoc(
        model, implColumns(includeFees), _WIDTHS, _TEXT_COLUMNS,
        variant=variant, feeSchedule=feeSchedule, feeLevel=feeLevel,
        includeFees=includeFees, proposalId=proposalId,
        customFeesBy=customFeesBy, customFeesAt=customFeesAt)

    portPages = paginate(portDoc, PAGE_BUDGET_PT)
    riskPages = paginate(riskDoc, PAGE_BUDGET_PT)
    implPages = paginate(_tableSlice(implDoc, headerRow), PAGE_BUDGET_PT,
                         headerRows=(headerRow,))
    assumPages = paginate(assumDoc, PAGE_BUDGET_PT, headerRows=(1, 2))
    items = [item for group in model.get('groups', []) for item in group['items']]

    basisLine = '{} · {}'.format(basis.currency, basis.hedging)
    names = ' against '.join(str(r['name']) for r in results)
    money = renderNumber(mandate.mandateSize, '$#,##0')
    footer = ('Proposal UID ' + proposalId if proposalId
              else 'Draft — no delivered proposal for this scenario')
    preamble = _preamblePairs(implDoc, headerRow)

    slides = []
    if cover:
        slides.append(('cover', None, None, None, None))
    for part, page in enumerate(portPages, 1):
        slides.append(('table', portDoc, page,
                       ('Strategic Asset Allocation',
                        _continued('{} — {}, per asset category with '
                                   'portfolio-level estimates'.format(basisLine, names),
                                   part, len(portPages)),
                        ['Category rows show the share of the total portfolio; asset '
                         'rows sum to their category. A category one portfolio holds '
                         'and another does not prints 0.0, so the columns compare row '
                         'for row.',
                         'Estimates derive from the long-term assumptions on the '
                         'final slide, at a 2.5% risk-free rate.']), None))
    for part, page in enumerate(riskPages, 1):
        slides.append(('table', riskDoc, page,
                       ('Risk Dashboard',
                        _continued('Factor-based risk analytics — stress periods and '
                                   'tail-loss measures, nominal and real, per portfolio',
                                   part, len(riskPages)),
                        ['Stress rows show predicted performance had each portfolio '
                         'been held through the episode; losses print red, gains '
                         'green — the workbook’s rule, resolved per cell here.',
                         'VaR and CVaR are stated as positive loss magnitudes at 99% '
                         'confidence, over 1 month, 1 year and 3 years.']), None))
    for part, page in enumerate(implPages, 1):
        slides.append(('table', implDoc, page,
                       ('Implemented Model',
                        _continued(' · '.join(preamble) if preamble else basisLine,
                                   part, len(implPages)),
                        ['Weights are rounded to 2dp by largest remainder across the '
                         'whole table, so the column closes on exactly 100.00%; '
                         'notionals derive from the printed weight on the {} '
                         'mandate.'.format(money),
                         'A notional below its product’s minimum would print red on '
                         'white; the export refuses to deliver while any remains. '
                         'Management fees resolve from the schedule and level named '
                         'above.']), None))
    if items:
        slides.append(('charts', None, None,
                       ('Composition of the Implemented Model',
                        'Share of allocation by product attribute — five views of '
                        'the same 100%',
                        ['Slice colours follow the page’s palette by the alphabetical '
                         'rule the screen uses; printed shares close on exactly 100.0.']),
                       items))
    for part, page in enumerate(assumPages, 1):
        slides.append(('table', assumDoc, page,
                       ('Long-Term Capital Market Assumptions',
                        _continued('Per-asset estimates behind the analytics — risk '
                                   'premia ranges, volatility and modelling windows',
                                   part, len(assumPages)),
                        ['Range ends are inked red and green, low — mean — high '
                         'across; the table shows only the assets the strategic '
                         'slides hold.',
                         'Estimates are a property of the {} universe recorded at '
                         'bake time; modelling windows are per asset, at a 2.5% '
                         'risk-free rate.'.format(basisLine)]), None))

    presentation = Presentation()
    presentation.slide_width = _in(SLIDE_W_IN)
    presentation.slide_height = _in(SLIDE_H_IN)
    blank = presentation.slide_layouts[6]
    total = len(slides)

    for number, (kind, doc, page, meta, payload) in enumerate(slides, 1):
        slide = presentation.slides.add_slide(blank)
        if kind == 'cover':
            head = _textbox(slide, MARGIN_IN, 2.35, SLIDE_W_IN - 2 * MARGIN_IN, 1.0)
            _para(head, 'Investment Proposal', 34, _HEAD_INK, bold=True, first=True)
            rule = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, _in(MARGIN_IN), _in(3.35),
                _in(3.2), _pt(3))
            rule.fill.solid()
            rule.fill.fore_color.rgb = RGBColor.from_string(_HEAD_INK)
            rule.line.fill.background()
            rule.shadow.inherit = False
            lines = _textbox(slide, MARGIN_IN, 3.62, SLIDE_W_IN - 2 * MARGIN_IN, 1.6)
            details = [basisLine + (' · ' + variant if variant else ''),
                       'Mandate {} · {}'.format(money, mandate.primaryPwa),
                       datetime.date.today().strftime('%d %B %Y')]
            for index, line in enumerate(details):
                _para(lines, line, 13, _SUB_INK, first=index == 0)
            page_ = _textbox(slide, SLIDE_W_IN - MARGIN_IN - 2.5,
                             SLIDE_H_IN - 0.58, 2.5, 0.5)
            _para(page_, footer, 8, _NOTE_INK, first=True, align=PP_ALIGN.RIGHT)
            _para(page_, 'Slide 1 of {}'.format(total), 8, _NOTE_INK,
                  align=PP_ALIGN.RIGHT)
            continue
        heading, subheading, footnotes = meta
        _chrome(slide, heading, subheading, footnotes, footer,
                'Slide {} of {}'.format(number, total))
        if kind == 'charts':
            _renderDonuts(slide, payload)
        else:
            _renderTable(slide, doc, page)

    presentation.core_properties.title = ('PMG Proposal {} - {} {}'.format(
        proposalId, basis.currency, basis.hedging) if proposalId
        else 'PMG Proposal - {} {}'.format(basis.currency, basis.hedging))
    presentation.core_properties.author = 'PMG Proposal Tool'
    presentation.core_properties.comments = (
        'Strategic allocation, risk dashboard, long-term estimates and the '
        'implemented model.')
    if proposalId:
        presentation.core_properties.identifier = proposalId
        presentation.core_properties.keywords = proposalId

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()
