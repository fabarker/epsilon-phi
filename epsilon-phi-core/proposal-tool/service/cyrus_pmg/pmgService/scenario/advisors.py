"""The Primary PWA directory - Excel-backed, as the reference adapter is
specified to be (spec 4.2 item 7). Production swaps the workbook read for a
database table behind the same two functions.

``advisors.xlsx`` sits beside this module: one sheet, columns ``name`` and
``office``. The display string the UI shows and stores is
``"{name} — {office}"``.

The two-character threshold is mirrored server-side: below it the search
returns nothing, matching the UI hint behaviour rather than dumping the
directory.
"""

from __future__ import annotations

import os
import threading

_XLSX = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'advisors.xlsx')
_lock = threading.Lock()
_directory = None


def _load() -> list:
    global _directory
    with _lock:
        if _directory is None:
            from openpyxl import load_workbook
            book = load_workbook(_XLSX, read_only=True)
            sheet = book.active
            rows = sheet.iter_rows(min_row=2, values_only=True)
            _directory = [
                {
                    'name': str(name).strip(),
                    'office': str(office).strip(),
                    'display': '{} — {}'.format(str(name).strip(), str(office).strip()),
                }
                for name, office in rows
                if name and office
            ]
            book.close()
    return _directory


def searchAdvisors(query: str, limit: int = 20) -> list:
    """Case-insensitive substring search over name, office and display."""
    query = (query or '').strip().lower()
    if len(query) < 2:
        return []
    hits = [
        advisor for advisor in _load()
        if query in advisor['display'].lower()
    ]
    return hits[:max(1, int(limit))]


def advisorExists(display: str) -> bool:
    """True when *display* is exactly a directory entry - free text is invalid."""
    display = (display or '').strip()
    return any(advisor['display'] == display for advisor in _load())
