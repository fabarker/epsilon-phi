"""Excel writers - the whole workbook, with no analytics library behind it.

Four sheets, and none of them needs ``Reporting`` or a ``SAAPortfolio``
(D67). Every figure on ``portfolios`` and ``risk_dashboard`` was already
carried by the resolved payloads the export passes in; ``assumptions`` reads
the per-asset estimates the bake now records beside them; ``Implementation``
was always the tool's own work. That is what lets the service be ported
somewhere the analytics library cannot go.

Three jobs live here:

* ``buildImplementationRows`` + ``writeImplementationSheet`` - the
  implementation model as a sheet (spec 14.4), with the rounding rule of spec
  8.4 applied exactly as on screen: every printed figure derives from the
  printed weight, so the workbook hand-reconciles against the page.
  ``buildImplementationRows`` is the single Python source of those numbers -
  the real adapter and the fixtures adapter both use it, and the test suite
  checks it against the page's JavaScript mirror.

* ``writePortfoliosSheet`` / ``writeRiskDashboardSheet`` /
  ``writeAssumptionsSheet`` - the three sheets the analytics library used to
  lay out, reproduced from the payloads cell for cell. Layout, fonts, fills,
  number formats, merges, row heights and the one conditional-formatting rule
  are the house report's own (spec 14.3), taken from a captured reference
  workbook rather than guessed at. Two conventions are deliberate and easy to
  break: asset rows carry weights x100 with format ``0.0`` while category and
  total rows carry fractions with ``0.0%``; and VaR and CVaR are stored
  NEGATED in the payload (the screen reads them as losses) but printed
  positive here.

* ``writeWorkbook`` - the four sheets in order, the only export path.
"""

from __future__ import annotations

import datetime
import io
import zipfile
from xml.etree import ElementTree

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from openpyxl.chart import DoughnutChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.formatting.rule import CellIsRule

from . import fees, rules
from . import portfolio_weights as pw
from .payloads import roundWeightsLargestRemainder
from .rules import sleeveCategory
from .types import PortfolioKey
from .sleeves import listSleeves

# The palette, the border kinds and every layout decision now live in
# sheetDoc (D119): this module is the Excel RENDERER. The aliases keep every
# name the tests and older callers reach for.
from . import sheetDoc
from .sheetDoc import (buildAssumptionsDoc, buildImplementationDoc,
                       buildPortfoliosDoc, buildRiskDoc,
                       _asDate, _categoryOrder, _weight)

_NAVY = sheetDoc.NAVY
_HEADER_NAVY = sheetDoc.HEADER_NAVY
_SUBHEAD_NAVY = sheetDoc.SUBHEAD_NAVY
_BAND = sheetDoc.BAND
_WHITE = sheetDoc.WHITE
_BREACH_FILL = sheetDoc.BREACH_FILL
_BREACH_INK = sheetDoc.BREACH_INK
_PREMIA_LOW = sheetDoc.PREMIA_LOW
_PREMIA_HIGH = sheetDoc.PREMIA_HIGH
_HEADING_INK = sheetDoc.HEADING_INK
_ASSUMPTION_WIDTHS = sheetDoc.ASSUMPTION_WIDTHS
_ASSUMPTION_COLUMNS = sheetDoc.ASSUMPTION_COLUMNS

# The composition doughnuts under the implementation table, and the palette
# the page draws them in (--cat-1..7). A slice takes its colour from where its
# name falls in the ALPHABETICAL order of that dimension's names, then the
# slices are ordered biggest-first to read - the same two steps the page's
# breakdown() does, so a chart in the workbook is the chart on the screen.
DONUT_DIMENSIONS = [('style', 'Style'), ('vehicle', 'Vehicle'), ('source', 'Source'),
                    ('liquidity', 'Liquidity'), ('exposureCurrency', 'Exposure currency')]
DONUT_PALETTE = ['2A78D6', 'EB6834', '1BAF7A', 'EDA100', 'E87BA4', '4A3AA7', 'E34948']
_CHART_SHEET = 'chartData'

# The columns of spec 9.3, less the two that left the sheet on request while
# staying on the screen: Ticker and Minimum Investment (D78, and see D74 on
# the screen and the sheet keeping their own header lists).
IMPL_COLUMNS = [
    'Categories & Asset Classes', 'Products', 'Allocation (%)',
    'Style', 'Vehicle', 'Share Class', 'Source', 'Liquidity', 'Exposure ccy',
    'Product Cost', 'Mgmt fee', 'Wtd fee (bp)', 'Notional',
]
#: The columns whose values are words rather than figures, left-aligned in the
#: body. Named rather than sliced by position, so a column leaving the list
#: cannot silently re-align its neighbours.
_TEXT_COLUMNS = ('Style', 'Vehicle', 'Share Class', 'Source', 'Liquidity', 'Exposure ccy')
# The two the sheet loses when the proposal excludes fees (D52). A proposal
# that does not show fees must not ship a sheet with empty columns and a
# header saying which schedule priced them: the columns go, and so do the fee
# rows above the header.
FEE_COLUMNS = ('Mgmt fee', 'Wtd fee (bp)')
# Ticker and Minimum Investment are not here: both left the SHEET on request
# (D78) while staying on the screen, in the catalogue and in the register.
_WIDTHS = {
    'Categories & Asset Classes': 34, 'Products': 32, 'Allocation (%)': 12,
    'Style': 9, 'Vehicle': 12, 'Share Class': 12, 'Source': 10, 'Liquidity': 11,
    'Exposure ccy': 12, 'Product Cost': 12, 'Mgmt fee': 10,
    'Wtd fee (bp)': 12, 'Notional': 14,
}


def implColumns(includeFees: bool = True) -> list:
    """The sheet's columns, in order, for a priced or an unpriced proposal."""
    if includeFees:
        return list(IMPL_COLUMNS)
    return [name for name in IMPL_COLUMNS if name not in FEE_COLUMNS]


def buildImplementationRows(baseResult: dict, sleevesMap: dict,
                            autoCategories, mandateSize: float,
                            variant: str = None, tacticalTilt: bool = False,
                            feeSchedule: str = None, feeLevel: str = None,
                            topAccountSize: float = None,
                            volPremium: bool = False,
                            currency: str = None,
                            customFees: dict = None) -> dict:
    """The implementation model's numbers, derived per spec 8.4.

    Returns {'groups': [...], 'total': {...}, 'complete': bool, 'priced': bool,
    'tier': {...} | None}. Each group is a category: its sleeve name (auto
    categories take their single library sleeve), product line items with
    printed weight, weighted fee in bp and notional, and group subtotals.
    Weights are rounded to 2dp by largest remainder across the WHOLE table so
    the column sums to exactly 100.00; notional and weighted fee derive from
    that printed weight.

    The management fee is resolved per product from *feeSchedule*, the tier
    that *topAccountSize* falls in, *feeLevel* and the product's fee group
    (D51). Without a schedule the model is unpriced: ``managementFee`` and
    ``wtdFeeBp`` are None on every item and in the total. The level defaults
    to the framework's prescribed one, as the store does; the schedule never
    defaults.
    """
    autoCategories = set(autoCategories or [])
    groups = []
    lineItems = []
    priced = feeSchedule is not None
    # A marginal schedule reads the MANDATE across the whole ladder and never
    # the top account size; a flat one reads the single tier the top account
    # size falls in. The tier is carried only where it means something (D83).
    marginal = priced and fees.isMarginal(feeSchedule)
    tier = fees.tierFor(topAccountSize) if priced and not marginal else None
    feeLevel = feeLevel or fees.DEFAULT_LEVEL
    # Under the custom level the ladder is not read at all: the one rate a
    # uniform schedule carries is the PWA's (D96), so that is the "blend".
    custom = priced and fees.isCustom(feeLevel)
    if marginal and custom:
        effectiveRate = fees.customRate(customFees, feeSchedule)
    elif marginal:
        effectiveRate = fees.effectiveRate(feeSchedule, mandateSize, feeLevel)
    else:
        effectiveRate = None

    # The implemented book, not the strategic one: with the tilt on, the
    # funding category is reduced and the tilt category appended (D50); with
    # the volatility premium on, the same category gives up a share of what is
    # left and Hybrid Fixed Income follows it (D53). The currency goes in
    # because the premium is forbidden outside USD and GBP, and the rule
    # belongs where the model is built, not only where the toggle is drawn.
    categories = rules.implementedCategories(baseResult['categories'],
                                             tacticalTilt, volPremium, currency)

    # Categories that share one sleeve share one LINE. Private Equity and
    # Other Private Assets are a single choice in the rail (D60) and a single
    # sleeve in the repository; showing them as two rows against the same
    # sleeve printed that sleeve twice and split a position that is bought
    # once. The combined weight is the sum of theirs, and the row takes the
    # group's own name. Driven off SLEEVE_GROUPS, so a second group needs no
    # code here.
    combined, order = {}, []
    for category in categories:
        key = sleeveCategory(category['name'])
        if key not in combined:
            combined[key] = {'name': key, 'weightPct': 0.0}
            order.append(key)
        combined[key]['weightPct'] += float(category['weightPct'])
    categories = [combined[key] for key in order]

    # the sleeves are resolved for the portfolio being implemented: a name
    # may hold several editions and only the one for this book is built (D89)
    baseKey = PortfolioKey.fromStr(baseResult['keyStr']) if baseResult.get('keyStr') else None
    for category in categories:
        name = category['name']
        catWeight = float(category['weightPct'])
        # grouped categories share one choice, stored under the group (D60)
        pickedUnder = sleeveCategory(name)
        if name in autoCategories:
            library = listSleeves(pickedUnder, variant, baseKey)
            sleeve = library[0] if library else None
            auto = True
        else:
            chosen = (sleevesMap or {}).get(pickedUnder)
            sleeve = None
            if chosen:
                sleeve = next((s for s in listSleeves(pickedUnder, variant, baseKey)
                               if s['name'] == chosen), None)
            auto = False
        items = []
        if sleeve:
            for product in sleeve['products']:
                item = dict(product)
                item['exactPct'] = catWeight * float(product['weight'])
                items.append(item)
                lineItems.append(item)
        groups.append({
            'category': name,
            'weightPct': catWeight,
            'sleeve': sleeve['name'] if sleeve else None,
            'auto': auto,
            'items': items,
        })

    complete = all(group['sleeve'] for group in groups) and bool(groups)

    if lineItems:
        if complete:
            printed = roundWeightsLargestRemainder([i['exactPct'] for i in lineItems])
        else:
            printed = [round(i['exactPct'], 2) for i in lineItems]
        for item, weight in zip(lineItems, printed):
            item['printedPct'] = weight
            item['notional'] = round(mandateSize * weight / 100.0 / 100.0) * 100.0
            if priced:
                item['managementFee'] = fees.productFee(
                    feeSchedule, feeLevel, item['feeGroup'],
                    topAccountSize=topAccountSize, mandateSize=mandateSize,
                    customFees=customFees)
                if item['managementFee'] is None:        # a custom row not yet set
                    item['wtdFeeBp'] = None
                else:
                    allIn = float(item['productCost']) + item['managementFee']
                    item['wtdFeeBp'] = allIn * weight    # percent x percent = bp
            else:
                item['managementFee'] = None
                item['wtdFeeBp'] = None
            # A position smaller than the product will accept is not a
            # position (item 3). Flagged per line here; the export refuses
            # while any survives, so the block cannot be walked past.
            minimum = item.get('minimumInvestment')
            item['belowMinimum'] = bool(minimum) and item['notional'] < float(minimum)

    # the rows the custom level has not priced, named so the export can say
    # which - one unpriced product makes the total unpriced, never smaller
    unpricedGroups = sorted({
        i['feeGroup'] if fees.byGroup(feeSchedule) else feeSchedule
        for i in lineItems if priced and i.get('managementFee') is None})
    total = {
        'weightPct': sum(i.get('printedPct', 0.0) for i in lineItems),
        'wtdFeeBp': (sum(i.get('wtdFeeBp') or 0.0 for i in lineItems)
                     if priced and not unpricedGroups else None),
        'notional': sum(i.get('notional', 0.0) for i in lineItems),
    }

    # Nothing that prints as zero earns a line (D68). Dropped AFTER the
    # rounding, so the largest-remainder pass still closes the column on
    # 100.00 exactly - a row worth 0.00 contributes nothing to that sum and
    # taking it out cannot move it. A category keeps its row while it has
    # weight, sleeve or no sleeve: an unimplemented category with an
    # allocation is the thing the page is asking a PWA to fix.
    for group in groups:
        group['items'] = [item for item in group['items']
                          if item.get('printedPct', 0.0) != 0]
    groups = [group for group in groups if group['weightPct'] != 0]

    breaches = [{'category': group['category'], 'name': item.get('name'),
                 'productId': item.get('productId'), 'notional': item['notional'],
                 'minimumInvestment': float(item['minimumInvestment'])}
                for group in groups for item in group['items']
                if item.get('belowMinimum')]

    # what the sheet's header states under the custom level: the row (or
    # rows) and the rate each carries, in the card's own order
    customRates = None
    if custom:
        customRates = [(row, fees.customRate(customFees, feeSchedule, row))
                       for row in fees.customRows(feeSchedule)]

    return {'groups': groups, 'total': total, 'complete': complete,
            'priced': priced, 'tier': tier, 'marginal': marginal,
            'effectiveRate': effectiveRate, 'breaches': breaches,
            'custom': custom, 'customRates': customRates,
            'unpricedGroups': unpricedGroups}


SHEET_PASSWORD = 'PA55WORD'
"""The workbook's sheet-protection password (D104).

Not a secret and not a security control: Excel's sheet protection is a
convention, the hash it stores is trivially stripped, and the file is a zip
besides. It is here so that unprotecting a delivered sheet is deliberate
rather than accidental, which is the whole of what this mechanism can do.
Nothing is encrypted, so nothing in the workbook is confidential by virtue
of it - the register's stored copy (D69) remains the record of what was
actually delivered.
"""


def protectSheet(sheet) -> None:
    """Lock the figures, leave the presentation alone (D98, D104).

    Every cell is already locked; protection is what makes that mean
    anything. The flags read backwards - True BLOCKS the action - so the
    three that are set False are the ones a reader is invited to use:
    formatting cells, column widths and row heights. Sorting and filtering
    are allowed too, because neither changes a value.

    The password is set on every sheet, including the hidden chart data, so
    the protection is uniform: a sheet left open would be the one a figure
    got edited on. openpyxl hashes it on assignment; the plaintext is never
    written to the file.
    """
    protection = sheet.protection
    protection.sheet = True          # the values are read only
    protection.password = SHEET_PASSWORD
    protection.formatCells = False   # fonts, fills, borders, number formats
    protection.formatColumns = False # column widths
    protection.formatRows = False    # row heights
    protection.sort = False
    protection.autoFilter = False


def _renderDoc(sheet, doc) -> None:
    """Transcribe one SheetDoc onto an openpyxl sheet - mechanically (D119).

    Every decision was made by the builder: this function maps neutral style
    attributes to openpyxl objects and holds no literal of its own. A None
    attribute means the builder did not touch it, so neither does this."""
    for rowIndex in sorted(doc.rows):
        rowSpec = doc.rows[rowIndex]
        if rowSpec.height is not None:
            sheet.row_dimensions[rowIndex].height = rowSpec.height
        for column in sorted(rowSpec.cells):
            spec = rowSpec.cells[column]
            cell = sheet.cell(row=rowIndex, column=column)
            if spec.value is not None:
                cell.value = spec.value
            if spec.fmt is not None:
                cell.number_format = spec.fmt
            if spec.font is not None:
                cell.font = Font(**spec.font)
            if spec.fill is not None:
                cell.fill = PatternFill('solid', fgColor=spec.fill)
            if spec.border is not None:
                cell.border = Border(**{
                    edge: (Side(style=style, color=colour) if colour
                           else Side(style=style))
                    for edge, (style, colour) in spec.border.items()})
            if spec.align is not None:
                cell.alignment = Alignment(**{
                    ('wrap_text' if key == 'wrap' else key): value
                    for key, value in spec.align.items()})
    for row1, col1, row2, col2 in doc.merges:
        sheet.merge_cells(start_row=row1, start_column=col1,
                          end_row=row2, end_column=col2)
    for letter, width in doc.widths.items():
        sheet.column_dimensions[letter].width = width
    if doc.freeze:
        sheet.freeze_panes = doc.freeze
    for col1, row1, col2, row2, negInk, posInk in doc.signRules:
        area = '{}{}:{}{}'.format(get_column_letter(col1), row1,
                                  get_column_letter(col2), row2)
        sheet.conditional_formatting.add(
            area, CellIsRule(operator='lessThan', formula=['0'],
                             font=Font(color=negInk)))
        sheet.conditional_formatting.add(
            area, CellIsRule(operator='greaterThan', formula=['0'],
                             font=Font(color=posInk)))


def roundSharesOneDp(exact) -> list:
    """Round shares to 1dp so they sum to exactly 100.0. Mirrors the page's
    roundSharesOneDp(); the two must agree or the export contradicts the UI."""
    units = [int(share * 10 + 1e-9) for share in exact]
    short = 1000 - sum(units)
    order = sorted(range(len(exact)),
                   key=lambda i: (-(exact[i] * 10 - units[i]), i))
    for n in range(max(short, 0)):
        if not order:
            break
        units[order[n % len(order)]] += 1
    return [unit / 10.0 for unit in units]


def donutBreakdown(items, key) -> list:
    """One doughnut's slices: share of what is attached, biggest first, each
    keeping the colour its name earns alphabetically. Mirrors the page's
    breakdown() exactly - see DONUT_PALETTE."""
    byValue, total = {}, 0.0
    for item in items:
        value = item.get(key)
        if value is None or value == '':
            value = '\u2014'
        byValue[value] = byValue.get(value, 0.0) + float(item.get('printedPct') or 0.0)
        total += float(item.get('printedPct') or 0.0)
    names = sorted(byValue)
    slot = {name: index % len(DONUT_PALETTE) for index, name in enumerate(names)}
    slices = [{'name': name, 'weight': byValue[name], 'slot': slot[name],
               'share': (byValue[name] / total) if total > 0 else 0.0}
              for name in names]
    slices.sort(key=lambda entry: (-entry['weight'], entry['name']))
    # printed shares close on 100.0 exactly, the way the page prints them
    printed = roundSharesOneDp([entry['share'] * 100 for entry in slices])
    for entry, pct in zip(slices, printed):
        entry['pct'] = pct
    return slices


def writeDonutCharts(book, sheet, model, firstRow: int) -> None:
    """The five composition doughnuts, under the implementation table.

    Native charts rather than pictures, so the figures stay live and the file
    stays small. Excel charts must read from cells, and those cells have no
    business on a sheet a client reads, so they go on a hidden sheet.
    """
    items = [item for group in model.get('groups', []) for item in group['items']]
    if not items:
        return
    data = book.create_sheet(_CHART_SHEET)
    data.sheet_state = 'hidden'

    sheet.cell(row=firstRow, column=1).value = 'Composition of the Implemented Model'
    sheet.cell(row=firstRow, column=1).font = Font(name='Calibri', size=12, bold=True,
                                                   color=_NAVY)
    sheet.cell(row=firstRow + 1, column=1).value = (
        'Share of allocation by product attribute.')
    sheet.cell(row=firstRow + 1, column=1).font = Font(name='Calibri', size=10,
                                                       color='5B6B7C')

    column = 1
    for index, (key, label) in enumerate(DONUT_DIMENSIONS):
        slices = donutBreakdown(items, key)
        if not slices:
            continue
        data.cell(row=1, column=column).value = label
        data.cell(row=1, column=column + 1).value = 'Share'
        for offset, entry in enumerate(slices, start=1):
            data.cell(row=1 + offset, column=column).value = entry['name']
            share = data.cell(row=1 + offset, column=column + 1)
            share.value = entry['pct'] / 100.0
            share.number_format = '0.0%'

        chart = DoughnutChart(holeSize=55)
        chart.title = label
        chart.height, chart.width = 7.4, 7.4
        chart.add_data(Reference(data, min_col=column + 1, min_row=1,
                                 max_row=1 + len(slices)), titles_from_data=True)
        chart.set_categories(Reference(data, min_col=column, min_row=2,
                                       max_row=1 + len(slices)))
        # one point per slice, in the palette the page uses
        series = chart.series[0]
        for offset, entry in enumerate(slices):
            point = DataPoint(idx=offset)
            point.graphicalProperties.solidFill = DONUT_PALETTE[entry['slot']]
            point.graphicalProperties.line.solidFill = 'FFFFFF'
            series.data_points.append(point)
        chart.dataLabels = None
        # laid out across the sheet, in the order the page shows them
        sheet.add_chart(chart, '{}{}'.format(
            get_column_letter(1 + index * 4), firstRow + 3))
        column += 2


def writeImplementationSheet(book, baseResult: dict, sleevesMap: dict,
                             autoCategories, mandateSize: float,
                             variant: str = None,
                             tacticalTilt: bool = False,
                             feeSchedule: str = None, feeLevel: str = None,
                             topAccountSize: float = None,
                             includeFees: bool = True,
                             volPremium: bool = False,
                             currency: str = None,
                             model: dict = None,
                             proposalId: str = None,
                             customFees: dict = None,
                             customFeesBy: str = None,
                             customFeesAt=None) -> None:
    """Append the implementation sheet: the columns of ``implColumns``, in
    order, grouped by category with subtotals and a grand total.

    Presentation is decided by ``sheetDoc.buildImplementationDoc`` (D119);
    what this function still owns is Excel: creating the sheet, rendering
    the Doc onto it, and the doughnut charts beneath the table. The variant,
    schedule, level and tier ride in the Doc's preamble for the reasons D29,
    D51, D52 and D96 give; ``buildImplementationDoc`` keeps those words.

    With *includeFees* false the proposal does not show fees at all (D52):
    the fee columns and fee header rows are absent, and the model is built
    unpriced whatever schedule the scenario happens to remember.
    """
    if not includeFees:
        feeSchedule = None
    columns = implColumns(includeFees)
    # A caller that has already built the model passes it in, so that what
    # the sheet prints and what the register records are the SAME model
    # rather than two builds a few microseconds apart (D69).
    if model is None:
        model = buildImplementationRows(baseResult, sleevesMap, autoCategories,
                                        mandateSize, variant, tacticalTilt,
                                        feeSchedule, feeLevel, topAccountSize,
                                        volPremium, currency, customFees)
    sheet = book.create_sheet('Implementation')
    doc, headerRow, totalRow = buildImplementationDoc(
        model, columns, _WIDTHS, _TEXT_COLUMNS,
        variant=variant, feeSchedule=feeSchedule, feeLevel=feeLevel,
        includeFees=includeFees, proposalId=proposalId,
        customFeesBy=customFeesBy, customFeesAt=customFeesAt)
    _renderDoc(sheet, doc)

    # the composition doughnuts, under the table the page draws them under
    writeDonutCharts(book, sheet, model, totalRow + 3)


# --------------------------------------------------------------------- #
# The three sheets the analytics library used to lay out (D67).
#
# Each is written from the resolved payloads alone. The layout is not
# invented: it was read cell by cell off a workbook the engine produced, and
# the test suite compares what these functions emit against captured golden
# files so a change here cannot drift away from the house report unnoticed.
# --------------------------------------------------------------------- #

def writePortfoliosSheet(book, results, engineParity: bool = False):
    """The strategic allocation, one column per portfolio - decided by
    ``sheetDoc.buildPortfoliosDoc`` (D119), rendered here."""
    sheet = book.create_sheet('portfolios')
    _renderDoc(sheet, buildPortfoliosDoc(results, engineParity))
    return sheet


def writeRiskDashboardSheet(book, results, engineParity: bool = False):
    """Risk, one pair of columns per portfolio - decided by
    ``sheetDoc.buildRiskDoc`` (D119), rendered here, the stress block's
    red/green rule included: the Doc declares it once and the renderer
    projects it as the same ``CellIsRule`` as ever."""
    sheet = book.create_sheet('risk_dashboard')
    _renderDoc(sheet, buildRiskDoc(results, engineParity))
    return sheet


def writeAssumptionsSheet(book, assets, results=None, engineParity: bool = False):
    """The long-term estimates behind the analytics, one row per asset (D67)
    - decided by ``sheetDoc.buildAssumptionsDoc`` (D119), rendered here."""
    sheet = book.create_sheet('assumptions')
    _renderDoc(sheet, buildAssumptionsDoc(assets, results, engineParity))
    return sheet


def writeWorkbook(basis, mandate, results, sleevesMap, autoCategories,
                  variant: str = None, tacticalTilt: bool = False,
                  feeSchedule: str = None, feeLevel: str = None,
                  includeFees: bool = True, volPremium: bool = False,
                  assets=None, engineParity: bool = False, model: dict = None,
                  proposalId: str = None, customFees: dict = None,
                  customFeesBy: str = None, customFeesAt=None) -> bytes:
    """The proposal workbook: four sheets, no analytics library (D67).

    ``portfolios``, ``risk_dashboard`` and ``assumptions`` reproduce what the
    engine's reporting object used to lay out; ``Implementation`` is the
    tool's own and is unchanged. *assets* is the per-asset estimate block for
    this basis, which the adapter supplies.

    *engineParity* restores the empty universe rows the library used to print
    (see ``_categoryOrder``). It exists for the golden test and should not be
    set by a caller producing a proposal.

    *proposalId* is the Proposal UID minted for this delivery (D75). It is
    written where a reader looks first (the Implementation sheet's first row),
    where a printed page shows it (every sheet's header) and where a program
    reads it (``dc:identifier``, see ``stampedProposalId``). Absent for a
    workbook that is not a delivery - the golden test, a unit test.
    """
    book = Workbook()
    book.remove(book.active)                      # the writers name their own

    writePortfoliosSheet(book, results, engineParity)
    writeRiskDashboardSheet(book, results, engineParity)
    writeAssumptionsSheet(book, assets, results, engineParity)
    writeImplementationSheet(book, results[0], sleevesMap, autoCategories,
                             mandate.mandateSize, variant, tacticalTilt,
                             feeSchedule, feeLevel, mandate.topAccountSize,
                             includeFees, volPremium, basis.currency,
                             model=model, proposalId=proposalId,
                             customFees=customFees, customFeesBy=customFeesBy,
                             customFeesAt=customFeesAt)

    # Enhancements that add nothing to the grid and cost nothing to read: a
    # coloured tab per sheet, a sensible print setup, and the proposal's own
    # identity in the file's properties.
    for name in book.sheetnames:
        sheet = book[name]
        # the figures are ours; the presentation is the reader's (D98). Every
        # sheet, the hidden chart data included - editing that would leave the
        # doughnuts disagreeing with the table they were drawn from.
        protectSheet(sheet)
        sheet.sheet_properties.tabColor = _NAVY
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_title_rows = '1:1'
        if proposalId:
            sheet.oddHeader.right.text = 'Proposal UID ' + proposalId
    book.properties.title = 'PMG Proposal - {} {}'.format(basis.currency, basis.hedging)
    book.properties.creator = 'PMG Proposal Tool'
    book.properties.description = (
        'Strategic allocation, risk dashboard, long-term estimates and the '
        'implemented model.')
    if proposalId:
        book.properties.title = 'PMG Proposal {} - {} {}'.format(
            proposalId, basis.currency, basis.hedging)
        book.properties.identifier = proposalId
        book.properties.keywords = proposalId

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def stampedProposalId(content: bytes):
    """The Proposal UID a workbook was written with, or None.

    Read from the file's own properties (``dc:identifier``) without opening a
    sheet, so it costs a zip entry rather than a load. None for bytes that are
    not a workbook, or a workbook written with no UID. The register compares
    this with the id it is about to record, and refuses a mismatch (D75).
    """
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            core = archive.read('docProps/core.xml')
    except (zipfile.BadZipFile, KeyError, ValueError):
        return None
    try:
        root = ElementTree.fromstring(core)
    except ElementTree.ParseError:
        return None
    node = root.find('{http://purl.org/dc/elements/1.1/}identifier')
    text = (node.text or '').strip() if node is not None else ''
    return text or None


#: The name this had while it was only the fixtures adapter's writer. It is
#: every adapter's writer now; the old name stays so nothing has to change at
#: once.
writeFixturesWorkbook = writeWorkbook
