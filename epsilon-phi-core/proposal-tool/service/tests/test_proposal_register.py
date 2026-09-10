"""The proposal register (D69).

Every delivered proposal, kept for ever: the two pictures, the sleeve pins,
and the workbook byte for byte. The store under test is the per-session
temporary database conftest points SCENARIO_REGISTER_DB at.
"""

import hashlib
import io
import json
import os
import re
import sqlite3

import pytest
from openpyxl import load_workbook

from cyrus_pmg.pmgService import dashboardRouter
from cyrus_pmg.pmgService.scenario import (assetEstimates, proposalRegister, rules,
                                           scenarioStore, sleeveRepo)
from cyrus_pmg.pmgService.scenario.types import BasisInput, MandateInput, ValidationError
from cyrus_pmg.pmgService.scenario.workbook import buildImplementationRows, writeWorkbook


def _caller(kerberos='alice'):
    """The identity a router dependency hands a handler.

    Both the mirror and the host pass a UserData now, so a test that calls a
    handler directly has to build one. Direct calls to the stores still pass
    the kerberos itself - that is what their TEXT columns hold.
    """
    from cyrus_pmg.pmgService.core import accessControl
    return accessControl.UserData(kerberos, accessControl._rolesFor(kerberos))

HERE = os.path.dirname(os.path.abspath(__file__))
GOLDEN = os.path.join(HERE, 'golden')


def _case():
    with open(os.path.join(GOLDEN, 'engine_USD_Hedged_2col.results.json'), encoding='utf-8') as fh:
        results = json.load(fh)
    with open(os.path.join(GOLDEN, 'engine_USD_Hedged_2col.impl.json'), encoding='utf-8') as fh:
        implementation = json.load(fh)
    return results, implementation


def _deliver(user='alice', createdBy='bob', scenarioId='sc_test', includeFees=True, results=None,
             recordAs=None, stampAs=None, filename=None):
    """Do what the export endpoint does: mint the UID, build the model once,
    write the workbook with both, name the file after the UID, record it.

    The three keyword arguments exist to break the one-UID invariant on
    purpose: *recordAs* records under another id, *stampAs* writes another
    id into the file ('' for none), *filename* names the file otherwise."""
    results0, implementation = _case()
    results = results or results0
    implementation = dict(implementation, includeFees=includeFees)
    basis = BasisInput(currency='USD', hedging='Hedged')
    mandate = MandateInput(topAccountSize=50_000_000.0, mandateSize=50_000_000.0,
                           primaryPwa='A. Castellanos — Madrid')
    model = buildImplementationRows(
        results[0], implementation['sleeves'], rules.AUTO_SLEEVE_CATEGORIES,
        mandate.mandateSize, implementation['variant'], implementation['tacticalTilt'],
        implementation['feeSchedule'] if includeFees else None, implementation['feeLevel'],
        mandate.topAccountSize, implementation['volPremium'], 'USD')
    proposalId = proposalRegister.newProposalId()
    content = writeWorkbook(
        basis, mandate, results, implementation['sleeves'], rules.AUTO_SLEEVE_CATEGORIES,
        implementation['variant'], implementation['tacticalTilt'],
        implementation['feeSchedule'], implementation['feeLevel'], includeFees,
        implementation['volPremium'], assets=assetEstimates.forSlice('USD', 'Hedged'),
        model=model, proposalId=proposalId if stampAs is None else stampAs)
    entry = proposalRegister.record(recordAs or proposalId, scenarioId, user, createdBy, basis,
                                    mandate, results, implementation, model, content,
                                    filename or rules.exportFilename(basis, proposalId))
    return entry, content, model


# --------------------------------------------------------------- recording ---

def test_a_delivered_proposal_is_recorded_whole():
    entry, content, model = _deliver()
    assert entry['proposalId'].startswith('pr_')
    assert entry['exportedBy'] == 'alice' and entry['createdBy'] == 'bob'
    assert entry['primaryPwa'] == 'A. Castellanos — Madrid'
    assert entry['topAccountSize'] == 50_000_000.0 and entry['mandateSize'] == 50_000_000.0
    assert entry['currency'] == 'USD' and entry['hedging'] == 'Hedged'
    assert entry['variant'] == 'PMG Multi-Asset Portfolio'
    assert entry['baseKey'] == 'USD|Moderate|Full|0'
    assert entry['workbookBytes'] == len(content)
    assert entry['workbookSha'] == hashlib.sha256(content).hexdigest()


def test_the_stored_workbook_is_byte_identical_to_the_one_delivered():
    """The whole reason to keep the file."""
    entry, content, _ = _deliver()
    kept = proposalRegister.workbook(entry['proposalId'])
    assert kept['bytes'] == content
    assert kept['name'] == entry['workbookName']
    assert kept['sha'] == hashlib.sha256(content).hexdigest()
    assert load_workbook(io.BytesIO(kept['bytes'])).sheetnames == [
        'portfolios', 'risk_dashboard', 'assumptions', 'Implementation', 'chartData']


def test_the_implemented_picture_matches_the_workbooks_implementation_sheet():
    """Picture two is what printed, row for row - not a recomputation."""
    entry, content, _ = _deliver()
    sheet = load_workbook(io.BytesIO(content))['Implementation']
    header = next(r for r in range(1, 12) if sheet.cell(row=r, column=1).value == 'Categories & Asset Classes')
    # the doughnut captions sit below the total now (item 9): read the table
    total = next(r for r in range(header, sheet.max_row + 1)
                 if sheet.cell(row=r, column=1).value == 'Total')
    printed = []
    for r in range(header + 1, total):
        label, product, weight = (sheet.cell(row=r, column=c).value for c in (1, 2, 3))
        if label and not str(label).startswith('  '):
            printed.append(('group', str(label).strip(), round(float(weight) * 100, 4)))
        elif product:
            printed.append(('item', str(product).strip(), round(float(weight) * 100, 4)))

    pictured = []
    for group in entry['implemented']:
        pictured.append(('group', group['category'], round(group['weightPct'], 4)))
        for item in group['items']:
            pictured.append(('item', item['name'], round(item['printedPct'], 4)))
    assert pictured == printed


def test_the_pictures_carry_no_analytics_and_no_fee_figures():
    entry, _, _ = _deliver()
    allocation = entry['allocation']
    assert set(allocation) == {'keyStr', 'name', 'header', 'categories'}
    for category in allocation['categories']:
        assert set(category) == {'name', 'weightPct', 'assets'}
    for group in entry['implemented']:
        for item in group['items']:
            for banned in ('managementFee', 'wtdFeeBp', 'feeGroup', 'productCost', 'exactPct'):
                assert banned not in item, banned
            assert item['name'] and 'printedPct' in item and 'notional' in item
    assert 'metrics' not in json.dumps(entry)


def test_only_the_proposal_portfolio_is_recorded_not_the_comparisons():
    """The columns a PWA compared against on screen are analysis. The register
    keeps the proposal - the base, the one the implementation is built on."""
    results, _ = _case()
    assert len(results) == 2, 'the fixture compares two portfolios'
    entry, _, _ = _deliver(scenarioId='sc_base_only')
    allocation = entry['allocation']
    assert isinstance(allocation, dict), 'one portfolio, not a list of columns'
    assert allocation['keyStr'] == results[0]['keyStr'] == entry['baseKey']
    assert allocation['name'] == results[0]['name']
    assert results[1]['name'] not in json.dumps(allocation)
    assert 'columnCount' not in entry and 'comparisons' not in entry


def test_the_pricing_plan_is_a_column_but_no_rate_is():
    """The plan name, without any figure: 'how was it priced', not 'what did it cost'."""
    priced, _, _ = _deliver(includeFees=True, scenarioId='sc_priced')
    assert priced['includeFees'] is True
    assert priced['feeSchedule'] == 'CASP' and priced['feeLevel'] == 'PMG Target'
    unpriced, _, _ = _deliver(includeFees=False, scenarioId='sc_unpriced')
    assert unpriced['includeFees'] is False
    assert unpriced['feeSchedule'] is None and unpriced['feeLevel'] is None
    assert 'tier' not in priced and 'wtdFeeBp' not in json.dumps(priced)


def test_every_sleeve_is_pinned_to_the_revision_it_was_at():
    entry, _, _ = _deliver()
    pins = {s['category']: s for s in entry['sleeves']}
    assert pins, 'a proposal with sleeves records its pins'
    for category, pin in pins.items():
        assert pin['sleeveId'] is not None and pin['revision'] is not None, category
        live = sleeveRepo.getSleeve(pin['sleeveId'])
        assert live['name'] == pin['sleeve']
        assert pin['revision'] == live['revisions'], 'pinned at the current revision'
        assert pin['moved'] is False and pin['nowRevision'] == pin['revision']


def test_the_register_sees_a_sleeve_move_after_the_proposal_was_made():
    """The one thing the register computes rather than stores."""
    entry, _, _ = _deliver(scenarioId='sc_drift')
    pin = next(s for s in entry['sleeves'] if s['sleeveId'] is not None)
    live = sleeveRepo.getSleeve(pin['sleeveId'])
    products = [{'productId': p['productId'], 'weight': p['weight']} for p in live['products']]
    sleeveRepo.updateSleeve(pin['sleeveId'], live['name'], products, note='moved after', user='eve')
    try:
        again = proposalRegister.getProposal(entry['proposalId'])
        moved = next(s for s in again['sleeves'] if s['sleeveId'] == pin['sleeveId'])
        assert moved['revision'] == pin['revision'], 'the pin does not move'
        assert moved['nowRevision'] == pin['revision'] + 1
        assert moved['moved'] is True
    finally:
        sleeveRepo.revertSleeve(pin['sleeveId'], pin['revision'], user='eve')


def test_re_exporting_the_same_scenario_is_a_second_proposal():
    first, _, _ = _deliver(scenarioId='sc_twice')
    second, _, _ = _deliver(scenarioId='sc_twice')
    assert first['proposalId'] != second['proposalId']
    assert (first['sequence'], second['sequence']) == (1, 2)
    assert first['scenarioId'] == second['scenarioId'] == 'sc_twice'


def test_nothing_to_record_is_refused():
    results, implementation = _case()
    basis = BasisInput(currency='USD', hedging='Hedged')
    mandate = MandateInput(50e6, 50e6, 'x')
    with pytest.raises(ValidationError):
        proposalRegister.record(proposalRegister.newProposalId(), 'sc_x', 'alice', 'bob', basis,
                                mandate, results, implementation, {'groups': []}, b'', 'x.xlsx')


def test_the_register_is_append_only_by_construction():
    """No update, no delete: the module simply has no such function."""
    names = {n for n in dir(proposalRegister) if not n.startswith('_')}
    for forbidden in ('update', 'delete', 'remove', 'purge', 'edit'):
        assert not [n for n in names if forbidden in n.lower()], forbidden
    import sqlite3
    conn = sqlite3.connect(proposalRegister.dbPath())
    columns = {r[1] for r in conn.execute('PRAGMA table_info(proposals)')}
    conn.close()
    assert 'updatedAt' not in columns and 'deletedAt' not in columns


# ----------------------------------------------------------------- reading ---

def test_listing_pages_newest_first_without_overlap_and_without_blobs():
    for i in range(5):
        _deliver(user='pager', scenarioId='sc_page%d' % i)
    page = proposalRegister.listProposals(exportedBy='pager', limit=2)
    assert len(page['entries']) == 2 and page['next']
    assert page['entries'][0]['exportedAt'] >= page['entries'][1]['exportedAt']
    assert 'workbook' not in page['entries'][0] and 'allocation' not in page['entries'][0]
    assert page['entries'][0]['workbookBytes'] > 10000, 'the size travels, the bytes do not'
    rest = proposalRegister.listProposals(exportedBy='pager', limit=10, before=page['next'])
    seen = {e['proposalId'] for e in page['entries']}
    assert not (seen & {e['proposalId'] for e in rest['entries']})
    assert page['total'] == 5


def test_filters_and_facets_count_what_the_other_filters_leave():
    _deliver(user='zed', scenarioId='sc_zed')
    only = proposalRegister.listProposals(exportedBy='zed', currency='USD')
    assert only['entries'] and all(e['exportedBy'] == 'zed' for e in only['entries'])
    assert 'zed' in only['facets']['exportedBy']
    assert 'alice' in only['facets']['exportedBy'], 'the person facet leaves the person out'
    assert 'USD' in only['facets']['currency']
    empty = proposalRegister.listProposals(currency='JPY')
    assert empty['entries'] == [] and empty['total'] == 0


def test_the_text_search_reaches_pwa_person_and_sleeve_names():
    entry, _, _ = _deliver(user='searchme', scenarioId='sc_search')
    sleeveName = entry['sleeves'][0]['sleeve']
    assert entry['proposalId'] in {e['proposalId'] for e in
                                   proposalRegister.listProposals(query='Castellanos', limit=500)['entries']}
    assert entry['proposalId'] in {e['proposalId'] for e in
                                   proposalRegister.listProposals(query='searchme')['entries']}
    assert entry['proposalId'] in {e['proposalId'] for e in
                                   proposalRegister.listProposals(query=sleeveName, limit=500)['entries']}
    assert proposalRegister.listProposals(query='zzz-nothing-zzz')['entries'] == []


def test_the_csv_export_carries_the_pins():
    entry, _, _ = _deliver(user='csvuser', scenarioId='sc_csv')
    rows = proposalRegister.exportRows(exportedBy='csvuser')
    assert rows and len(rows[0]) == len(proposalRegister.EXPORT_COLUMNS)
    row = dict(zip(proposalRegister.EXPORT_COLUMNS, rows[0]))
    assert row['ProposalId'] == entry['proposalId']
    assert '@r' in row['Sleeves'] and row['WorkbookSha'] == entry['workbookSha']


def test_describe_counts_proposals_and_bytes():
    d = proposalRegister.describe()
    assert d['proposals'] >= 1 and d['workbookBytes'] > 10000
    assert d['schemaVersion'] == proposalRegister.SCHEMA_VERSION and d['earliest']


# --------------------------------------------------------------- endpoints ---

def test_the_scenario_store_records_who_started_it():
    state = scenarioStore.createScenario(MandateInput(50e6, 50e6, 'A. Castellanos — Madrid'),
                                         BasisInput(currency='USD', hedging='Hedged'),
                                         createdBy='starter')
    assert state['createdBy'] == 'starter'


def test_the_endpoints_answer_as_the_panel_expects(monkeypatch):
    monkeypatch.setenv('PMG_ALLOWED_KERBEROS', 'alice')
    monkeypatch.setenv('PMG_ADMIN_KERBEROS', 'alice')
    entry, content, _ = _deliver(user='alice', scenarioId='sc_endpoint')

    page = dashboardRouter.listRegisterProposals(exportedBy='alice', caller=_caller('alice'))
    assert entry['proposalId'] in {e['proposalId'] for e in page['entries']}

    one = dashboardRouter.getRegisterProposal(entry['proposalId'], caller=_caller('alice'))['proposal']
    assert one['allocation'] and one['implemented'] and one['sleeves']
    assert 'workbook' not in one

    file = dashboardRouter.downloadRegisterWorkbook(entry['proposalId'], caller=_caller('alice'))
    assert file.body == content
    assert file.headers['x-workbook-sha256'] == entry['workbookSha']
    assert 'attachment' in file.headers['content-disposition']

    missing = dashboardRouter.getRegisterProposal('pr_nope', caller=_caller('alice'))
    assert missing.status_code == 404

    csv_ = dashboardRouter.exportRegisterProposals(exportedBy='alice', caller=_caller('alice'))
    assert csv_.body.decode('utf-8').splitlines()[0] == ','.join(proposalRegister.EXPORT_COLUMNS)

    body = dashboardRouter.getRepository(caller=_caller('alice'))
    assert body['register']['proposals'] >= 1


# ---------------------------------------------------------------- the UID ---
# One UID per finished proposal (D75): minted before the file is written,
# visible in the file and in its name, and the register's primary key.

def _stampedIn(content):
    """Everywhere the UID shows in a delivered file."""
    book = load_workbook(io.BytesIO(content))
    sheet = book['Implementation']
    return {
        'a1': (sheet.cell(row=1, column=1).value, sheet.cell(row=1, column=2).value),
        'printHeaders': {name: book[name].oddHeader.right.text for name in book.sheetnames},
        'identifier': book.properties.identifier,
        'title': book.properties.title,
    }


def test_one_uid_is_in_the_row_the_file_and_the_filename():
    from cyrus_pmg.pmgService.scenario.workbook import stampedProposalId
    entry, content, _ = _deliver()
    uid = entry['proposalId']
    assert proposalRegister.isProposalId(uid)
    seen = _stampedIn(content)
    assert seen['a1'] == ('Proposal UID', uid), 'the first cell a reader meets'
    assert seen['identifier'] == uid and uid in seen['title']
    assert set(seen['printHeaders'].values()) == {'Proposal UID ' + uid}, 'every printed page'
    assert stampedProposalId(content) == uid, 'what the register reads back'
    assert entry['workbookName'].startswith('PMG_Scenario_USD_Hedged_')
    assert entry['workbookName'].endswith('_{}.xlsx'.format(uid))
    assert proposalRegister.workbook(uid)['name'] == entry['workbookName']


def test_the_uid_row_sits_above_the_implementation_type():
    entry, content, _ = _deliver()
    sheet = load_workbook(io.BytesIO(content))['Implementation']
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][:2] == ('Proposal UID', entry['proposalId'])
    assert rows[1][0] == 'Implementation Type'
    header = next(i for i, r in enumerate(rows) if r[0] == 'Categories & Asset Classes')
    assert rows[header - 1][0] is None, 'the blank row still parts the preamble from the table'
    assert sheet.freeze_panes == 'A{}'.format(header + 2), 'the pane freezes under the moved header'


def test_a_workbook_with_no_uid_is_not_a_delivery():
    """A file written without the UID cannot be recorded under one."""
    from cyrus_pmg.pmgService.scenario.workbook import stampedProposalId
    with pytest.raises(ValidationError) as caught:
        _deliver(stampAs='')
    assert caught.value.field == 'proposalId'
    assert stampedProposalId(b'not a workbook') is None


def test_a_workbook_stamped_with_another_uid_is_not_recorded():
    """The register checks rather than trusts: the file must say what the row says."""
    with pytest.raises(ValidationError) as caught:
        _deliver(stampAs=proposalRegister.newProposalId())
    assert caught.value.field == 'proposalId'


def test_a_filename_without_the_uid_is_not_recorded():
    with pytest.raises(ValidationError) as caught:
        _deliver(filename='PMG_Scenario_USD_Hedged_2026-09-05.xlsx')
    assert caught.value.field == 'proposalId'


def test_something_that_is_not_a_uid_is_refused():
    with pytest.raises(ValidationError):
        _deliver(recordAs='PR-123', stampAs='PR-123', filename='PMG_PR-123.xlsx')


def test_a_uid_is_never_reused():
    """Primary key: a second row for an id is an error, not an overwrite."""
    entry, _, _ = _deliver()
    uid = entry['proposalId']
    with pytest.raises(sqlite3.IntegrityError):
        _deliver(recordAs=uid, stampAs=uid, filename=entry['workbookName'])
    assert proposalRegister.getProposal(uid)['sequence'] == entry['sequence']


def test_the_export_filename_carries_the_uid_last():
    name = rules.exportFilename(BasisInput(currency='GBP', hedging='Hedged'), 'pr_0123456789ab')
    assert re.fullmatch(r'PMG_Scenario_GBP_Hedged_\d{4}-\d{2}-\d{2}_pr_0123456789ab\.xlsx', name)


def test_minted_uids_have_one_shape_and_do_not_repeat():
    minted = {proposalRegister.newProposalId() for _ in range(2000)}
    assert len(minted) == 2000
    assert all(proposalRegister.isProposalId(uid) for uid in minted)
    assert not proposalRegister.isProposalId('pr_0123456789AB'), 'lower case only'
    assert not proposalRegister.isProposalId('sc_0123456789ab')
