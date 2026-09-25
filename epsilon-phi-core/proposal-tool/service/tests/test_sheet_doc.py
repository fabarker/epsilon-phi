"""Phase 1 of the PowerPoint export plan (D119): the report as data.

Every layout and styling decision now lives in ``sheetDoc``'s builders, and
``workbook.py`` is the Excel renderer. These tests pin the properties the
split exists for; the golden workbook test proves the split moved nothing.
"""

import io
import json
import os

from openpyxl import load_workbook

from cyrus_pmg.pmgService.scenario import assetEstimates, rules, sheetDoc
from cyrus_pmg.pmgService.scenario.workbook import (
    _TEXT_COLUMNS, _WIDTHS, buildImplementationRows, implColumns)

HERE = os.path.dirname(os.path.abspath(__file__))
GOLDEN = os.path.join(HERE, 'golden')


def _goldenCase():
    with open(os.path.join(GOLDEN, 'engine_USD_Hedged_2col.results.json'),
              encoding='utf-8') as handle:
        results = json.load(handle)
    with open(os.path.join(GOLDEN, 'engine_USD_Hedged_2col.impl.json'),
              encoding='utf-8') as handle:
        implementation = json.load(handle)
    return results, implementation


def _dump(doc):
    """A Doc as plain data, so two builds can be compared for identity."""
    return {
        'name': doc.name,
        'rows': {row: {'height': spec.height,
                       'cells': {col: (cell.value, cell.fmt, cell.font,
                                       cell.fill, cell.border, cell.align)
                                 for col, cell in spec.cells.items()}}
                 for row, spec in doc.rows.items()},
        'merges': doc.merges, 'widths': doc.widths, 'freeze': doc.freeze,
        'signRules': doc.signRules, 'sections': sorted(doc.sections),
    }


def _docs():
    results, implementation = _goldenCase()
    model = buildImplementationRows(
        results[0], implementation['sleeves'], rules.AUTO_SLEEVE_CATEGORIES,
        25_000_000, implementation['variant'], implementation['tacticalTilt'],
        implementation['feeSchedule'], implementation['feeLevel'],
        10_000_000, implementation['volPremium'], 'USD')
    impl, headerRow, totalRow = sheetDoc.buildImplementationDoc(
        model, implColumns(True), _WIDTHS, _TEXT_COLUMNS,
        variant=implementation['variant'],
        feeSchedule=implementation['feeSchedule'],
        feeLevel=implementation['feeLevel'],
        includeFees=True, proposalId='pr_0123456789ab')
    return {
        'portfolios': sheetDoc.buildPortfoliosDoc(results),
        'risk_dashboard': sheetDoc.buildRiskDoc(results),
        'assumptions': sheetDoc.buildAssumptionsDoc(
            assetEstimates.forSlice('USD', 'Hedged'), results),
        'Implementation': impl,
    }, headerRow, totalRow


def test_the_builders_import_no_file_format():
    """The one hard rule (C1): a style in the Doc is plain data. The module
    must not import openpyxl - or any other file format - however its
    comments talk about them. Read from the AST, so prose cannot trip it
    and an import cannot hide in it."""
    import ast
    path = os.path.join(HERE, '..', 'cyrus_pmg', 'pmgService', 'scenario', 'sheetDoc.py')
    with open(path, encoding='utf-8') as handle:
        tree = ast.parse(handle.read())
    imported = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            imported.update(alias.name for alias in node.names)
        elif isinstance(node, ast.ImportFrom):
            module = node.module or ''
            imported.add(module)
            imported.update('{}.{}'.format(module, alias.name) for alias in node.names)
    for banned in ('openpyxl', 'pptx', 'xlsxwriter'):
        assert not any(name == banned or name.startswith(banned + '.')
                       for name in imported), (banned, sorted(imported))
    # and nothing sideways: the exact import surface, pinned. A new import
    # here is a design event, and this list is where it gets reviewed.
    assert imported == {'__future__', '__future__.annotations',
                        'datetime', 'decimal',
                        'decimal.Decimal', 'decimal.ROUND_HALF_UP',
                        '', '.fees', '.portfolio_weights'}, sorted(imported)


def test_the_excel_writers_hold_no_styling_literal_of_their_own():
    """The four sheet writers are wrappers: create the sheet, render the Doc,
    add the Excel-only extras. A Font or Fill appearing in one again would
    mean a styling decision the deck cannot see (D119)."""
    path = os.path.join(HERE, '..', 'cyrus_pmg', 'pmgService', 'scenario', 'workbook.py')
    with open(path, encoding='utf-8') as handle:
        source = handle.read()
    wrappers = source[source.index('def writeImplementationSheet'):
                      source.index('def writeWorkbook')]
    for construct in ('Font(', 'PatternFill(', 'Border(', 'Alignment(',
                      'Side(', 'CellIsRule('):
        assert construct not in wrappers, construct


def test_the_builders_are_deterministic():
    """The same payloads produce the same Doc, attribute for attribute -
    the property that lets two renderers trust one build."""
    first, _, _ = _docs()
    second, _, _ = _docs()
    for name in first:
        assert _dump(first[name]) == _dump(second[name]), name


def test_the_docs_carry_the_sheets_shape():
    """A few load-bearing facts a renderer relies on, read from the Docs
    rather than from a rendered file."""
    docs, headerRow, totalRow = _docs()

    impl = docs['Implementation']
    header = [impl.cell(headerRow, col).value
              for col in range(1, len(implColumns(True)) + 1)]
    assert header == implColumns(True)
    assert impl.cell(1, 1).value == 'Proposal UID'
    assert impl.cell(totalRow, 1).value == 'Total'
    assert impl.freeze == 'A' + str(headerRow + 1)
    assert impl.sections, 'category bands mark the page-break rows'

    risk = docs['risk_dashboard']
    assert len(risk.signRules) == 1, 'the stress rule is declared once'
    (col1, row1, col2, row2, neg, pos) = risk.signRules[0]
    assert (neg, pos) == (sheetDoc.STRESS_NEG, sheetDoc.STRESS_POS)
    assert risk.cell(row1 - 1, 1).value == 'Predicted Performance Over Stress Periods'
    assert risk.merges, 'the nominal/real pairs merge on the summary rows'

    port = docs['portfolios']
    labels = [row.cells[1].value for row in port.rows.values()
              if 1 in row.cells and row.cells[1].value]
    assert 'TOTAL' in labels and 'Estimated Mean Return' in labels

    assum = docs['assumptions']
    assert assum.cell(1, 2).value == 'Long-Term Estimates'
    assert (2, 2, 2, 4) in [tuple(m) for m in assum.merges]


def test_a_doc_renders_the_same_sheet_twice():
    """Rendering is mechanical: the same Doc onto two sheets produces
    identical cells - the renderer holds no state and no opinion."""
    from openpyxl import Workbook
    from cyrus_pmg.pmgService.scenario.workbook import _renderDoc
    docs, _, _ = _docs()
    doc = docs['portfolios']
    book = Workbook()
    one = book.create_sheet('one')
    two = book.create_sheet('two')
    _renderDoc(one, doc)
    _renderDoc(two, doc)
    for row in range(1, one.max_row + 1):
        for col in range(1, one.max_column + 1):
            a, b = one.cell(row=row, column=col), two.cell(row=row, column=col)
            assert (a.value, a.number_format, a.font.name, a.font.bold,
                    a.fill.fgColor.rgb if a.fill.patternType else None) == \
                   (b.value, b.number_format, b.font.name, b.font.bold,
                    b.fill.fgColor.rgb if b.fill.patternType else None)


# --------------------------------------------------------------------- #
# Phase 2 (D120): renderNumber and paginate.
# --------------------------------------------------------------------- #

import datetime

import pytest


def test_render_number_displays_what_excel_displays():
    """One expectation per format the builders use, values hand-checked
    against Excel's display of the same cell."""
    rn = sheetDoc.renderNumber
    # percents scale in Decimal, so 0.2405 x 100 is 24.05, not 24.049999...
    assert rn(0.2405, '0.00%') == '24.05%'
    assert rn(0.34, '0.0%') == '34.0%'
    assert rn(-0.177, '0.0%') == '-17.7%'
    assert rn(1.0, '0%') == '100%'
    assert rn(0.005996386355446893, '0.0%') == '0.6%'
    # the sign rides the value: a tiny negative keeps its minus, as in Excel
    assert rn(-0.0004, '0.0%') == '-0.0%'
    # plain decimals round halves away from zero - Excel's rounding
    assert rn(8.36, '0.0') == '8.4'
    assert rn(34.0, '0.0') == '34.0'
    assert rn(0.125, '0.00') == '0.13'
    assert rn(0.4755554160115003, '0.00') == '0.48'
    # money: grouped, rounded to the unit, sign before the $
    assert rn(25000000.0, '$#,##0') == '$25,000,000'
    assert rn(787500.0, '$#,##0') == '$787,500'
    assert rn(1202500.5, '$#,##0') == '$1,202,501'
    assert rn(-1234, '$#,##0') == '-$1,234'
    # dates, locale-proof
    assert rn(datetime.datetime(1973, 2, 28), 'mmm-yy') == 'Feb-73'
    assert rn(datetime.datetime(2025, 4, 30), 'mmm-yy') == 'Apr-25'
    # General: what the sheets actually put under it
    assert rn(0, 'General') == '0'
    assert rn(0.0, 'General') == '0'
    assert rn(None, '0.0%') == ''
    # text ignores a numeric format, as in Excel
    assert rn('Total', '$#,##0') == 'Total'
    with pytest.raises(ValueError):
        rn(1.0, '#,##0.00')


def test_every_format_a_builder_uses_is_renderable():
    """The catalogue is closed: a builder cannot introduce a number format
    the deck does not know how to display. Walks every cell of the golden
    Docs and renders it."""
    docs, _, _ = _docs()
    seen = set()
    for doc in docs.values():
        for row in doc.rows.values():
            for cell in row.cells.values():
                if cell.fmt is not None:
                    assert cell.fmt in sheetDoc.NUMBER_FORMATS, (doc.name, cell.fmt)
                    seen.add(cell.fmt)
                sheetDoc.renderNumber(cell.value, cell.fmt)   # must not raise
    assert '0.00%' in seen and '$#,##0' in seen and 'mmm-yy' in seen


def _syntheticDoc(rows, sections, heights=None):
    doc = sheetDoc.SheetDoc('test')
    for index in rows:
        doc.cell(index, 1).value = 'r%d' % index
        if heights and index in heights:
            doc.row(index).height = heights[index]
    doc.sections.update(sections)
    return doc


def test_paginate_breaks_only_at_section_marks():
    """Budget for three body rows a page; sections at 2, 5 and 8 - the
    breaks land exactly on them, and the header repeats on every page."""
    doc = _syntheticDoc(range(1, 11), sections={2, 5, 8})
    pages = sheetDoc.paginate(doc, budgetPt=60.0)             # 15 header + 3 rows
    assert [p.body for p in pages] == [[2, 3, 4], [5, 6, 7], [8, 9, 10]]
    assert all(p.header == [1] for p in pages)


def test_paginate_walks_back_to_the_last_section_mark():
    """Room for three rows but the section starts mid-page: the break moves
    back so the section head starts the next page with its lines."""
    doc = _syntheticDoc(range(1, 8), sections={2, 4})
    pages = sheetDoc.paginate(doc, budgetPt=60.0)
    assert [p.body for p in pages] == [[2, 3], [4, 5, 6], [7]]


def test_paginate_splits_hard_when_a_section_outgrows_the_budget():
    """No mark inside the window: the stretch splits where it must rather
    than overflowing the canvas, and no row is lost or repeated."""
    doc = _syntheticDoc(range(1, 12), sections={2})
    pages = sheetDoc.paginate(doc, budgetPt=60.0)
    flat = [index for page in pages for index in page.body]
    assert flat == list(range(2, 12))
    assert all(len(p.body) <= 3 for p in pages)


def test_paginate_counts_heights_not_rows():
    """A 49-point header row costs what it costs: the budget is points, the
    way a canvas actually fills."""
    doc = _syntheticDoc(range(1, 6), sections={2}, heights={2: 49.0})
    pages = sheetDoc.paginate(doc, budgetPt=80.0)             # 15 + 49 + 15 = 79
    assert pages[0].body == [2, 3]
    assert pages[1].body == [4, 5]


def test_paginate_is_pure_and_covers_every_row_once():
    """The real Docs: nothing mutated, every body row on exactly one page,
    reading order preserved - with a one-page degenerate case."""
    docs, headerRow, _ = _docs()
    impl = docs['Implementation']
    before = _dump(impl)
    headers = tuple(range(1, headerRow + 1))
    pages = sheetDoc.paginate(impl, budgetPt=220.0, headerRows=headers)
    assert _dump(impl) == before, 'paginate must not touch the Doc'
    flat = [index for page in pages for index in page.body]
    assert flat == [index for index in sorted(impl.rows) if index > headerRow]
    assert len(pages) > 1
    # a break lands on a category band wherever one was available
    assert all(page.body[0] in impl.sections for page in pages[1:])
    whole = sheetDoc.paginate(impl, budgetPt=1e6, headerRows=headers)
    assert len(whole) == 1 and whole[0].body == flat
