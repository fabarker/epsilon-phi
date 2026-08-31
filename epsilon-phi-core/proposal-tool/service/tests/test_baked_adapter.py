"""Tier 1: the bake store and the adapter that serves it.

Baked with the fixtures port so these run with no database, exercising the
store format, the lookup, the miss/fallback behaviour and the promise that
matters — a COLD process (empty caches, nothing warmed) answers immediately.
"""

import json
import os
import sys
import time

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))

from cyrus_pmg.pmgService.scenario import bake, rules
from cyrus_pmg.pmgService.scenario.bakedAdapter import BakedScenarioPort
from cyrus_pmg.pmgService.scenario.fixturesAdapter import FixturesScenarioPort
from cyrus_pmg.pmgService.scenario.types import (
    AnalyticsError, BasisInput, MandateInput, PortfolioKey)

BASIS = BasisInput(currency='USD', hedging='Hedged')


@pytest.fixture(scope='module')
def store(tmp_path_factory):
    directory = str(tmp_path_factory.mktemp('baked'))
    bake.bakeSlice(FixturesScenarioPort(), 'USD', 'Hedged', directory, verbose=False)
    return directory


def test_slice_covers_the_whole_availability_set(store):
    payloads = json.load(open(bake.slicePath(store, 'USD', 'Hedged')))
    expected = rules.availability(BASIS, None)
    assert len(expected) == 68
    assert sorted(payloads) == sorted(expected)


def test_manifest_records_provenance_and_coverage(store):
    manifest = json.load(open(os.path.join(store, 'manifest.json')))
    assert manifest['portfoliosBaked'] == 68
    assert manifest['currencies'] == ['USD']
    entry = manifest['slices']['USD_Hedged.json']
    assert entry['available'] == 68 and entry['baked'] == 68
    assert entry['failures'] == {}
    assert entry['describe']['adapter'] == 'fixtures'


def test_baked_payloads_match_the_source_port(store):
    port = BakedScenarioPort(storeDirectory=store)
    live = FixturesScenarioPort()
    for keyStr in rules.availability(BASIS, None)[:6]:
        key = PortfolioKey.fromStr(keyStr)
        assert port.resolve_portfolio(BASIS, key) == live.resolve_portfolio(BASIS, key)


def test_cold_process_resolves_immediately(store):
    """The point of the whole exercise: no warm-up, no analytics, ~1ms."""
    keys = [PortfolioKey.fromStr(k) for k in rules.availability(BASIS, None)[:2]]
    port = BakedScenarioPort(storeDirectory=store)          # nothing loaded yet
    t0 = time.time()
    first = port.resolve_portfolio(BASIS, keys[0])          # includes the slice read
    firstMs = (time.time() - t0) * 1000
    t0 = time.time()
    port.resolve_portfolio(BASIS, keys[1])
    secondMs = (time.time() - t0) * 1000
    assert first['metrics']['volatilityPct'] > 0
    assert firstMs < 250, 'cold slice load took {:.0f}ms'.format(firstMs)
    assert secondMs < 5, 'warm lookup took {:.1f}ms'.format(secondMs)


def test_schema_assembly_is_cheap(store):
    """get_schema is fetched on open and on every basis/mandate change."""
    port = BakedScenarioPort(storeDirectory=store)
    port.get_schema(BASIS, None)                            # prime the memo
    t0 = time.time()
    for _ in range(20):
        port.get_schema(BASIS, MandateInput(48.5e6, 26e6, 'M. Aldridge — Zurich'))
    perCallMs = (time.time() - t0) * 1000 / 20
    assert perCallMs < 2, 'get_schema took {:.1f}ms per call'.format(perCallMs)


def test_miss_without_delegate_raises_analytics_error(store):
    port = BakedScenarioPort(storeDirectory=store)
    with pytest.raises(AnalyticsError):
        port.resolve_portfolio(BasisInput(currency='USD', hedging='Unhedged'),
                               PortfolioKey('Core', True, False, 'Mod'))


def test_miss_falls_through_to_the_delegate(store):
    port = BakedScenarioPort(storeDirectory=store, delegate=FixturesScenarioPort())
    unbaked = BasisInput(currency='USD', hedging='Unhedged')
    key = PortfolioKey('Core', True, False, 'Mod')
    result = port.resolve_portfolio(unbaked, key)
    # against the naming rule, not a frozen string: the point of this test is
    # that the miss reached the delegate, not what the delegate calls things
    assert result['name'] == rules.portfolioName(unbaked, key)


def test_schema_and_library_need_no_analytics(store):
    port = BakedScenarioPort(storeDirectory=store)
    schema = port.get_schema(BASIS, MandateInput(48.5e6, 26e6, 'M. Aldridge — Zurich'))
    assert len(schema['availability']) == 68
    assert schema['categories'][0] == 'Investment Grade Fixed Income'
    assert schema['dataInfo']['adapter'] == 'baked'
    assert '68 portfolios baked' in schema['dataInfo']['dataversion']
    from cyrus_pmg.pmgService.scenario.sleeves import VARIANTS
    assert port.list_sleeves('Public Equity', BASIS, VARIANTS[0])
    assert port.search_advisors('ald')


def test_export_without_a_delegate_uses_the_payload_writer(store, tmp_path):
    from openpyxl import load_workbook
    port = BakedScenarioPort(storeDirectory=store)
    key = PortfolioKey('Core', True, False, 'Mod')
    result = port.resolve_portfolio(BASIS, key)
    from cyrus_pmg.pmgService.scenario.sleeves import VARIANTS, listSleeves
    chosen = {c['name']: listSleeves(c['name'], VARIANTS[0])[0]['name']
              for c in result['categories']
              if c['name'] not in rules.AUTO_SLEEVE_CATEGORIES}
    mandate = MandateInput(48.5e6, 26e6, 'M. Aldridge — Zurich')
    path = tmp_path / 'baked.xlsx'
    path.write_bytes(port.build_export(
        BASIS, mandate, [result], {'sleeves': chosen, 'variant': VARIANTS[0]}))
    book = load_workbook(path)
    assert book.sheetnames == ['Portfolios', 'Risk Dashboard', 'Implementation']
    rows = list(book['Implementation'].iter_rows(values_only=True))
    weights = [r[2] for r in rows
               if r[2] is not None and r[0] and str(r[0]).startswith('  ')]
    assert abs(sum(weights) * 100 - 100.0) < 1e-9


def test_bake_is_resumable_and_records_failures(tmp_path):
    class Flaky(FixturesScenarioPort):
        def resolve_portfolio(self, basis, key):
            if key.riskLevel == 'Agg':
                raise AnalyticsError('simulated engine failure')
            return super().resolve_portfolio(basis, key)

    directory = str(tmp_path)
    entry = bake.bakeSlice(Flaky(), 'USD', 'Hedged', directory, verbose=False)
    assert entry['failures'] and all('simulated engine failure' in m
                                     for m in entry['failures'].values())
    assert entry['baked'] == entry['available'] - len(entry['failures'])

    # a resumed bake with a healthy port fills only the gaps and keeps the rest
    before = json.load(open(bake.slicePath(directory, 'USD', 'Hedged')))
    entry2 = bake.bakeSlice(FixturesScenarioPort(), 'USD', 'Hedged', directory,
                            verbose=False)
    after = json.load(open(bake.slicePath(directory, 'USD', 'Hedged')))
    assert entry2['resumedWith'] == len(before)
    assert entry2['baked'] == entry2['available'] and entry2['failures'] == {}
    for keyStr, payload in before.items():
        assert after[keyStr] == payload
