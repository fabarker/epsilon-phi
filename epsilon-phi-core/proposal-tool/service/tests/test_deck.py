"""Phase 3 of the PowerPoint export plan (D121): the deck renderer.

The same Docs the Excel renderer transcribes, rendered as slides - and held
to them cell for cell: every table cell's text must equal renderNumber of
the Doc's value under the Doc's format, the sign rule must resolve to the
CellIsRule's inks, the band fills and black rules must survive the trip,
and the doughnuts must carry the page's palette with the workbook's hole.
"""

import io
import json
import os

import pytest
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn

from cyrus_pmg.pmgService.scenario import assetEstimates, pptWriter, rules, sheetDoc
from cyrus_pmg.pmgService.scenario.pptWriter import writeDeck
from cyrus_pmg.pmgService.scenario.types import BasisInput, MandateInput
from cyrus_pmg.pmgService.scenario.workbook import (
    DONUT_PALETTE, buildImplementationRows, donutBreakdown, implColumns,
    stampedProposalId, _TEXT_COLUMNS, _WIDTHS)

HERE = os.path.dirname(os.path.abspath(__file__))
GOLDEN = os.path.join(HERE, 'golden')
UID = 'pr_0123456789ab'
BASIS = BasisInput(currency='USD', hedging='Hedged')
MANDATE = MandateInput(topAccountSize=10_000_000.0, mandateSize=25_000_000.0,
                       primaryPwa='J. Mercer')


def _goldenCase():
    with open(os.path.join(GOLDEN, 'engine_USD_Hedged_2col.results.json'),
              encoding='utf-8') as handle:
        results = json.load(handle)
    with open(os.path.join(GOLDEN, 'engine_USD_Hedged_2col.impl.json'),
              encoding='utf-8') as handle:
        implementation = json.load(handle)
    return results, implementation


def _model(results, implementation, includeFees=True):
    return buildImplementationRows(
        results[0], implementation['sleeves'], rules.AUTO_SLEEVE_CATEGORIES,
        MANDATE.mandateSize, implementation['variant'],
        implementation['tacticalTilt'],
        implementation['feeSchedule'] if includeFees else None,
        implementation['feeLevel'], MANDATE.topAccountSize,
        implementation['volPremium'], 'USD')


def _deck(proposalId=UID, includeFees=True, cover=True):
    results, implementation = _goldenCase()
    return writeDeck(BASIS, MANDATE, results, implementation['sleeves'],
                     rules.AUTO_SLEEVE_CATEGORIES, implementation['variant'],
                     implementation['tacticalTilt'],
                     implementation['feeSchedule'], implementation['feeLevel'],
                     includeFees, implementation['volPremium'],
                     assets=assetEstimates.forSlice('USD', 'Hedged'),
                     proposalId=proposalId, cover=cover)


def _pages():
    """The pages the deck should hold, computed the way writeDeck computes
    them - so the test states the mapping rather than trusting it."""
    results, implementation = _goldenCase()
    model = _model(results, implementation)
    port = sheetDoc.buildPortfoliosDoc(results)
    risk = sheetDoc.buildRiskDoc(results)
    assum = sheetDoc.buildAssumptionsDoc(
        assetEstimates.forSlice('USD', 'Hedged'), results)
    impl, headerRow, _ = sheetDoc.buildImplementationDoc(
        model, implColumns(True), _WIDTHS, _TEXT_COLUMNS,
        variant=implementation['variant'],
        feeSchedule=implementation['feeSchedule'],
        feeLevel=implementation['feeLevel'], includeFees=True,
        proposalId=UID)
    budget = pptWriter.PAGE_BUDGET_PT
    return [
        (port, sheetDoc.paginate(port, budget)),
        (risk, sheetDoc.paginate(risk, budget)),
        (impl, sheetDoc.paginate(pptWriter._tableSlice(impl, headerRow),
                                 budget, headerRows=(headerRow,))),
        (assum, sheetDoc.paginate(assum, budget, headerRows=(1, 2))),
    ]


def _table(slide):
    tables = [shape.table for shape in slide.shapes if shape.has_table]
    assert len(tables) == 1, 'one table per table slide'
    return tables[0]


def _normal(text):
    lines = str(text).split('\n')
    return '\n'.join([lines[0]] + [line.strip() for line in lines[1:]])


def test_the_deck_is_the_docs_cell_for_cell():
    """The heart of parity: walk every page of every Doc against its slide's
    table, and every cell's text must be renderNumber of the Doc's value
    under the Doc's format - the same decision both files render."""
    prs = Presentation(io.BytesIO(_deck()))
    sections = _pages()
    slideAt = 1                                   # slide 0 is the cover
    for doc, pages in sections:
        for page in pages:
            table = _table(prs.slides[slideAt])
            rows = page.header + page.body
            columns = pptWriter._docColumns(doc)
            assert (len(table.rows), len(table.columns)) == (len(rows), columns)
            for tableRow, rowIndex in enumerate(rows):
                spec = doc.rows.get(rowIndex)
                for column in range(1, columns + 1):
                    cell = (spec.cells.get(column) if spec is not None else None)
                    expected = ('' if cell is None
                                else sheetDoc.renderNumber(cell.value, cell.fmt))
                    actual = table.cell(tableRow, column - 1).text
                    assert _normal(actual) == _normal(expected), (
                        doc.name, rowIndex, column)
            slideAt += 1
        if doc.name == 'Implementation':
            slideAt += 1                          # the composition charts
    assert slideAt == len(prs.slides)


def test_the_stress_rule_resolves_to_the_workbooks_inks():
    """Red below zero, green above, base ink outside the rule's range -
    resolved from the raw value exactly as the CellIsRule would."""
    prs = Presentation(io.BytesIO(_deck()))
    riskSlide = prs.slides[2]                     # cover, portfolios, risk
    table = _table(riskSlide)
    inks = {}
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    inks.setdefault(run.text, str(run.font.color.rgb))
    assert inks['-17.7%'] == '9C0006', 'a loss in the stress block is red'
    assert inks['10.7%'] == '006100', 'a gain in the stress block is green'
    # VaR is positive but OUTSIDE the rule's range: base ink, not green
    assert inks['17.4%'] == '000000'
    assert inks['0.48'] == '000000', 'Sharpe sits outside the rule too'


def test_the_implementation_slide_keeps_the_bands_and_the_rules():
    """The category band's fill, the total's black rules, and no styling
    invented: left and right edges are explicit no-lines."""
    prs = Presentation(io.BytesIO(_deck()))
    table = _table(prs.slides[3])
    bandCell = totalCell = None
    for row in table.rows:
        cells = list(row.cells)
        if cells[0].text == 'Public Equity':
            bandCell = cells[0]
        if cells[0].text == 'Total':
            totalCell = cells[0]
    assert bandCell is not None and totalCell is not None
    assert str(bandCell.fill.fore_color.rgb) == 'D3DDEA'
    tcPr = totalCell._tc.get_or_add_tcPr()
    for edge, lit in (('a:lnT', True), ('a:lnB', True), ('a:lnL', False)):
        lines = tcPr.findall(qn(edge))
        assert len(lines) == 1, edge
        fills = lines[0].findall(qn('a:solidFill'))
        if lit:
            assert fills and fills[0].find(qn('a:srgbClr')).get('val') == '000000'
        else:
            assert lines[0].find(qn('a:noFill')) is not None


def test_the_doughnuts_are_native_with_the_pages_palette_and_the_hole():
    """Five doughnut charts, slice colours by the alphabetical palette rule,
    the workbook's 55% hole written where python-pptx has no property."""
    results, implementation = _goldenCase()
    model = _model(results, implementation)
    items = [item for group in model['groups'] for item in group['items']]
    prs = Presentation(io.BytesIO(_deck()))
    chartSlide = prs.slides[4]                    # after the implementation
    charts = [shape.chart for shape in chartSlide.shapes if shape.has_chart]
    assert len(charts) == 5
    for chart in charts:
        assert chart.chart_type == XL_CHART_TYPE.DOUGHNUT
        holes = chart._chartSpace.findall('.//' + qn('c:holeSize'))
        assert holes and all(hole.get('val') == '55' for hole in holes)
    styleSlices = donutBreakdown(items, 'style')
    points = charts[0].series[0].points
    got = [str(points[index].format.fill.fore_color.rgb)
           for index in range(len(styleSlices))]
    assert got == [DONUT_PALETTE[entry['slot']] for entry in styleSlices]


def test_the_deck_is_stamped_like_a_workbook_and_a_draft_is_not():
    """dc:identifier carries the UID, so stampedProposalId reads a deck
    exactly as it reads a workbook; a deck with no delivered proposal says
    Draft in every footer and stamps nothing."""
    deck = _deck()
    assert stampedProposalId(deck) == UID
    prs = Presentation(io.BytesIO(deck))
    assert prs.core_properties.title == 'PMG Proposal {} - USD Hedged'.format(UID)

    draft = _deck(proposalId=None)
    assert stampedProposalId(draft) is None
    drafted = Presentation(io.BytesIO(draft))
    texts = ' '.join(shape.text_frame.text for slide in drafted.slides
                     for shape in slide.shapes if shape.has_text_frame)
    assert 'Draft — no delivered proposal for this scenario' in texts
    assert UID not in texts


def test_the_fee_columns_leave_the_deck_with_the_fees():
    """includeFees=False: the two fee columns are absent from the slide as
    they are absent from the sheet, because implColumns drives both."""
    prs = Presentation(io.BytesIO(_deck(includeFees=False)))
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_table:
                continue
            header = [cell.text for cell in next(iter(shape.table.rows)).cells]
            if 'Products' in header:
                assert 'Mgmt fee' not in header
                assert 'Wtd fee (bp)' not in header
                assert 'Product Cost' in header, 'product cost stays (D52)'
                return
    pytest.fail('no implementation table found')


def test_the_cover_is_a_flag():
    withCover = Presentation(io.BytesIO(_deck()))
    without = Presentation(io.BytesIO(_deck(cover=False)))
    assert len(withCover.slides) == len(without.slides) + 1
    assert not any(shape.has_table for shape in withCover.slides[0].shapes)
    assert any(shape.has_table for shape in without.slides[0].shapes)
